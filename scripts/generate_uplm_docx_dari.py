# -*- coding: utf-8 -*-
"""
Build U_P_L_M_Dari.docx — full Dari monograph for NAMS (Exit.pdf structure).
RTL; B Nazanin / Tahoma; APA-style tables; same assets as generate_uplm_docx.py.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
MONO = next(p for p in ROOT.iterdir() if p.is_dir() and (p / "exit.pdf").exists())
ASSETS = MONO / "_uplm_assets"
OUT = MONO / "U_P_L_M_Dari.docx"

# B Nazanin preferred for Dari theses; Tahoma is installed on this Windows host.
FONT_PRIMARY = "B Nazanin"
FONT_FALLBACK = "Tahoma"
BODY_PT = 12
HEAD_PT = 14
LINE = 480  # double spacing (APA)


def _resolve_font() -> str:
    fonts_dir = Path(r"C:\Windows\Fonts")
    candidates = [
        ("B Nazanin", ["BNazanin.ttf", "B Nazanin.ttf", "bnazanin.ttf", "BNazanin.TTF"]),
        ("Tahoma", ["tahoma.ttf", "TAHOMA.TTF"]),
        ("Arial", ["arial.ttf", "ARIAL.TTF"]),
    ]
    for name, files in candidates:
        if any((fonts_dir / f).exists() for f in files):
            return name
    return FONT_FALLBACK


FONT = _resolve_font()


def set_rtl(p):
    pPr = p._element.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(OxmlElement("w:bidi"))


def set_run(run, size=BODY_PT, bold=False, italic=False, font=None):
    f = font or FONT
    run.font.name = f
    run._element.rPr.rFonts.set(qn("w:ascii"), f)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), f)
    run._element.rPr.rFonts.set(qn("w:cs"), f)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), f)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_para(p, size=BODY_PT, bold=False, align="right", space_after=8, first_indent=True):
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    set_rtl(p)
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if first_indent and align in ("justify", "right"):
        pf.first_line_indent = Inches(0.5)
    else:
        pf.first_line_indent = Inches(0)
    for r in p.runs:
        set_run(r, size=size, bold=bold)


def add_p(doc, text, *, size=BODY_PT, bold=False, align="right", first_indent=True, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    style_para(p, size=size, bold=bold, align=align, space_after=space_after, first_indent=first_indent)
    return p


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=HEAD_PT, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(10)
    pf.first_line_indent = Inches(0)
    p.style = doc.styles["Heading %d" % min(level, 3)]
    for r in p.runs:
        set_run(r, size=HEAD_PT, bold=True)
    return p


def page_break(doc):
    doc.add_page_break()


def add_caption_table(doc, number, title):
    add_p(doc, f"جدول {number}", size=BODY_PT, bold=True, align="right", first_indent=False, space_after=0)
    add_p(doc, title, size=BODY_PT, bold=False, align="right", first_indent=False, space_after=6)


def add_caption_figure(doc, number, title):
    add_p(doc, f"شکل {number}", size=BODY_PT, bold=True, align="right", first_indent=False, space_after=0)
    add_p(doc, title, size=BODY_PT, bold=False, align="right", first_indent=False, space_after=6)
    p = doc.add_paragraph()
    run = p.add_run("یادداشت. از مصنوعات پروژه سیستم مدیریت اکادمی نکته (NAMS) و داده‌های ارزیابی تهیه شده است.")
    set_run(run, size=10, italic=True)
    style_para(p, size=10, align="right", first_indent=False, space_after=10)


def set_cell_text(cell, text, bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    set_run(run, size=size, bold=bold)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_rtl(p)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Inches(0)


def _set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        if val is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(val))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def apa_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for j, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], h, bold=True, size=11)
        _set_cell_borders(table.rows[0].cells[j], top=12, bottom=12, left=None, right=None)
    for i, row in enumerate(rows):
        is_last = i == len(rows) - 1
        for j, val in enumerate(row):
            set_cell_text(table.rows[i + 1].cells[j], val, bold=False, size=11)
            if is_last:
                _set_cell_borders(table.rows[i + 1].cells[j], top=None, bottom=12, left=None, right=None)
            else:
                _set_cell_borders(table.rows[i + 1].cells[j], top=None, bottom=None, left=None, right=None)
    add_p(doc, "", first_indent=False, space_after=4)
    return table


def add_image(doc, path: Path, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    return p


def figure_block(doc, path, number, title, before_paras, after_paras):
    for t in before_paras:
        add_p(doc, t)
    add_image(doc, path)
    add_caption_figure(doc, number, title)
    for t in after_paras:
        add_p(doc, t)


def table_block(doc, number, title, headers, rows, before_paras, after_paras):
    for t in before_paras:
        add_p(doc, t)
    add_caption_table(doc, number, title)
    apa_table(doc, headers, rows)
    for t in after_paras:
        add_p(doc, t)


def configure_styles(doc: Document):
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.size = Pt(HEAD_PT if "Heading" in style_name else BODY_PT)
        st.font.bold = "Heading" in style_name
        if st._element.rPr is not None and st._element.rPr.rFonts is not None:
            st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
            st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            st._element.rPr.rFonts.set(qn("w:cs"), FONT)
            st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def build_front_matter(doc: Document):
    add_p(doc, "پوهنتون کابل", size=16, bold=True, align="center", first_indent=False, space_after=12)
    add_p(doc, "پوهنځی کمپیوتر ساینس", size=15, bold=True, align="center", first_indent=False, space_after=12)
    add_p(doc, "دیپارتمنت انجینری نرم‌افزار", size=14, bold=True, align="center", first_indent=False, space_after=36)
    add_p(doc, "سیستم مدیریت اکادمی نکته (NAMS)", size=18, bold=True, align="center", first_indent=False, space_after=24)
    add_p(doc, "مونوگراف نهایی پروژه", size=15, bold=True, align="center", first_indent=False, space_after=36)
    add_p(doc, "محمد میوند عثمانی  |  محمد خالد ساحل", size=BODY_PT, align="center", first_indent=False, space_after=12)
    add_p(doc, "استاد رهنما: استاد عبدالواحد سمدزی", size=BODY_PT, align="center", first_indent=False, space_after=48)
    page_break(doc)

    add_h(doc, "چکیده", 1)
    add_p(
        doc,
        "اکادمی‌های خصوصی در افغانستان، از جمله اکادمی نکته، اغلب با ثبت‌های دستی، فایل‌های پراکنده، "
        "تأخیر در راپوردهی، ضعف ارتباطات و نبود شفافیت مالی روبه‌رو هستند. این مونوگراف سیستم مدیریت "
        "اکادمی نکته (NAMS) را معرفی می‌کند؛ پلتفرم یکپارچه ویب–API–موبایل برای ثبت‌نام شاگردان، "
        "مدیریت صنف و مضمون، حاضری، امتحانات، فیس، معاش، مالیه دولتی و داشبوردهای نقش‌محور. "
        "نیازمندی‌ها از طریق پرسشنامه ۲۰ اکادمی (۱۸ مورد نیاز به امکانات نمونه دوم، ۲ تازه‌تأسیس "
        "بدون نظر مشخص)، تحلیل جریان کاری و Agile Scrum در سه نمونه اولیه جمع‌آوری شد. "
        "نمونه اول عملیات پایه را بدون امتحان هوش مصنوعی، لینک Google Meet و حاضری خودکار والدین "
        "تحویل داد؛ نمونه دوم پس از تأیید نظرسنجی این امکانات را افزود؛ نمونه سوم گزارش مالیه "
        "فصل‌وار و سالانه دولتی را بر اساس درخواست اکادمی نکته پیاده‌سازی کرد. "
        "نتایج نشان می‌دهد NAMS می‌تواند کار دستی را کاهش دهد، شفافیت مالی را تقویت کند و "
        "پذیرش بالای استفاده را در نمونه هدفمند ۲۰ اکادمی نشان دهد (Creswell & Creswell, 2018; "
        "Pressman & Maxim, 2020).",
    )
    page_break(doc)


def build_chapter1(doc: Document):
    add_h(doc, "فصل اول: مقدمه", 1)

    add_h(doc, "۱.۱ پس‌زمینه", 2)
    add_p(
        doc,
        "تحول دیجیتالی برای نهادهای آموزشی به یک ضرورت تبدیل شده است؛ زیرا مدیریت دقیق معلومات، "
        "سرعت در امور اداری، شفافیت مالی و ارتباطات منظم بدون سیستم دیجیتالی دشوار است. "
        "در بسیاری از اکادمی‌های خصوصی افغانستان، ثبت‌نام، تقسیم اوقات، فیس، نمرات و حاضری هنوز "
        "به شکل کاغذی یا در فایل‌های Excel جداگانه انجام می‌شود. این روش در ابتدا ساده به نظر "
        "می‌رسد، اما با افزایش تعداد شاگردان، استادان و تراکنش‌های مالی، خطا، تکرار کار و "
        "تأخیر در تصمیم‌گیری را افزایش می‌دهد (Aldowah et al., 2019; Romero & Ventura, 2020).",
    )
    add_p(
        doc,
        "سیستم مدیریت اکادمی نکته (NAMS) به عنوان پلتفرم یکپارچه طراحی شده است تا عملیات آموزشی، "
        "اداری، ارتباطی و مالی را در یک محیط واحد مدیریت کند. این سیستم ثبت شاگرد، صنف، مضمون، "
        "حاضری، امتحان، نتیجه، پیام، فیس، معاش، مالیه و داشبورد را پشتیبانی می‌کند و با "
        "پروپوزل پروژه که Node.js، Express.js، MongoDB و React.js را پیشنهاد کرده، همخوان است.",
    )
    add_p(
        doc,
        "از منظر انجینری نرم‌افزار، اهمیت پروژه در تبدیل یک مشکل واقعی اکادمی به محصول قابل "
        "استفاده است. ادبیات تأکید می‌کند که سیستم موفق باید بر نیازمندی‌های واضح، معماری "
        "قابل نگهداری و تست ساخته شود (Pressman & Maxim, 2020; Sommerville, 2016).",
    )

    add_h(doc, "۱.۲ بیان مسئله", 2)
    add_p(
        doc,
        "اکادمی نکته و اکادمی‌های مشابه به دلیل نبود سیستم مرکزی دیجیتالی با مشکلات اداری و "
        "آموزشی روبه‌رو هستند. عملیات مهم مانند ثبت شاگرد، حاضری، نمرات، تقسیم اوقات، فیس، "
        "معاش و راپوردهی یا دستی انجام می‌شود یا در فایل‌های جدا نگهداری می‌گردد. "
        "این پراکندگی ناسازگاری معلومات، دوباره‌کاری و کاهش شفافیت را به دنبال دارد.",
    )
    add_p(
        doc,
        "مسئله دوم ارتباطات است. شاگرد، استاد، والدین و اداره به پلتفرم مشترک برای اعلان‌ها، "
        "تقسیم اوقات، نتایج و پیام‌ها نیاز دارند. نبود چنین پلتفرمی باعث تأخیر یا از دست "
        "رفتن معلومات می‌شود (DeLone & McLean, 2003; Venkatesh et al., 2003).",
    )
    add_p(
        doc,
        "مسئله سوم نبود نظارت علمی مبتنی بر داده است. پروپوزل بر تحلیل هوش مصنوعی و پیش‌بینی "
        "عملکرد تأکید دارد؛ در این پروژه تمرکز اصلی بر مدیریت و ساختار داده است، اما مدل "
        "داده برای توسعه آینده AI آماده شده است (Siemens, 2013).",
    )

    add_h(doc, "۱.۳ اهداف تحقیق", 2)
    add_h(doc, "۱.۳.۱ هدف عمومی", 3)
    add_p(
        doc,
        "طراحی، توسعه و پیاده‌سازی سیستم یکپارچه ویب–موبایل برای اکادمی نکته که عملیات "
        "آموزشی، اداری، ارتباطی و مالی را خودکار سازد و زمینه تحلیل‌های هوش مصنوعی را "
        "آماده کند.",
    )
    add_h(doc, "۱.۳.۲ اهداف اختصاصی", 3)
    table_block(
        doc,
        "۱.۱",
        "همخوانی سوالات تحقیق، اهداف و شواهد پروژه",
        ["سوال تحقیق", "هدف همخوان", "شواهد در پروژه"],
        [
            ["چگونه سیستم یکپارچه کارایی عملیاتی را بهبود می‌دهد؟", "ساخت پلتفرم مرکزی NAMS", "معماری، ماژول‌ها، داشبورد"],
            ["کدام ناکارایی‌ها با دیجیتالی‌سازی حل می‌شوند؟", "تحلیل جریان کار و پرسشنامه", "فصل ۳ و ۴"],
            ["تحلیل AI چگونه نظارت علمی را تقویت می‌کند؟", "ذخیره‌سازی ساختارمند داده", "حاضری، نتایج، امتحان"],
            ["بهترین روش ارتباط و شفافیت مالی چیست؟", "پیام، فیس، معاش، مالیه", "ماژول‌های مالی Prototype 1–3"],
            ["چگونه Node.js و React پلتفرم پایدار می‌سازند؟", "معماری لایه‌ای", "شکل ۳.۱۶ و جدول ۳.۴"],
        ],
        [
            "جدول ۱.۱ نشان می‌دهد هر سوال پروپوزل به هدف و بخش عملی پروژه وصل شده است.",
            "این ردیابی برای دفاع علمی و ارزیابی استاد راهنما ضروری است.",
        ],
        [
            "سوالات تحقیق در فصل پنج با شواهد پرسشنامه و نمونه‌های اولیه پاسخ داده می‌شوند.",
        ],
    )

    add_h(doc, "۱.۴ سوالات تحقیق", 2)
    add_p(
        doc,
        "سوالات تحقیق اصلی پروژه عبارت‌اند از: (۱) چگونه یک سیستم آنلاین یکپارچه "
        "کارایی عملیاتی و یادگیری شخصی‌سازی‌شده را بهبود می‌دهد؟ (۲) کدام "
        "ناکارایی‌های فعلی با دیجیتالی‌سازی برطرف می‌شوند؟ (۳) تحلیل‌های مبتنی "
        "بر AI چگونه نظارت علمی را تقویت می‌کنند؟ (۴) بهترین روش برای ارتباطات "
        "و شفافیت مالی چیست؟ (۵) Node.js، MongoDB و React چگونه پلتفرم پایدار "
        "می‌سازند؟ جدول ۱.۱ هر سوال را به هدف و شواهد پیاده‌سازی وصل می‌کند.",
    )
    add_p(
        doc,
        "این سوالات در فصل سوم از طریق mixed-methods و در فصل پنج از طریق "
        "ماتریس بحث (جدول ۵.۱) پاسخ داده می‌شوند (Hevner et al., 2004).",
    )

    add_h(doc, "۱.۵ محدوده تحقیق", 2)
    add_p(
        doc,
        "محدوده شامل سیستم ویب با API، ثبت شاگرد و استاد، صنف، مضمون، تقسیم اوقات، حاضری، "
        "امتحان، فیس، معاش، پیام، داشبورد، چندزبانه (دری، پشتو، انگلیسی) و سه نمونه اولیه "
        "است. کاربران هدف: سوپر ادمین، اداره، استاد، شاگرد، والدین، مالک و مدیر شعبه.",
    )
    add_p(
        doc,
        "خارج از محدوده: ادغام ملی در سطح وزارت، حضور بیومتریک، و مدل‌های پیش‌بینی ML "
        "پیشرفته که به عنوان کار آینده در نظر گرفته شده‌اند. ارزیابی پرسشنامه بر ۲۰ اکادمی "
        "هدفمند است، نه نماینده آماری کل کشور (Creswell & Creswell, 2018).",
    )

    add_h(doc, "۱.۶ اهمیت تحقیق", 2)
    add_p(
        doc,
        "اهمیت عملی: کاهش کار دستی، افزایش دقت معلومات و شفافیت مالی. اهمیت علمی: "
        "کاربرد انجینری نرم‌افزار، نیازسنجی مختلط و Agile Scrum در پروژه واقعی. "
        "اهمیت اجتماعی: کمک به دیجیتالی‌سازی سکتور آموزش خصوصی افغانستان.",
    )

    add_h(doc, "۱.۷ سازماندهی گزارش", 2)
    add_p(
        doc,
        "فصل اول مقدمه، مسئله، اهداف، محدوده و اهمیت را بیان می‌کند. فصل دوم ادبیات "
        "تحقیق را مرور می‌کند. فصل سوم میتودولوژی، Agile Scrum، سه نمونه اولیه، طراحی "
        "و پیاده‌سازی را شرح می‌دهد. فصل چهارم نتایج پرسشنامه را ارائه می‌کند. "
        "فصل پنج بحث و فصل شش نتیجه‌گیری است. منابع و ضمیمه‌ها در پایان آمده‌اند.",
    )
    page_break(doc)


def build_chapter2(doc: Document):
    add_h(doc, "فصل دوم: مرور ادبیات تحقیق", 1)

    add_h(doc, "۲.۱ مفاهیم و تعاریف", 2)
    add_p(
        doc,
        "سیستم مدیریت اکادمی یک سیستم معلوماتی آموزشی است که داده‌های عملیاتی، علمی، "
        "مالی و ارتباطی را یکجا مدیریت می‌کند. مدل موفقیت DeLone و McLean کیفیت سیستم، "
        "کیفیت معلومات، استفاده، رضایت و منفعت را ارزیابی می‌کند (DeLone & McLean, 2003).",
    )
    add_p(
        doc,
        "Learning Analytics و Educational Data Mining با پروپوزل AI پروژه مرتبط‌اند. "
        "ذخیره منظم حاضری، ارزیابی و فیس زمینه تحلیل آینده را فراهم می‌سازد "
        "(Aldowah et al., 2019; Romero & Ventura, 2020).",
    )
    add_p(
        doc,
        "مهندسی نیازمندی‌ها بر اساس ISO/IEC/IEEE 29148:2018 نیازمندی‌های واضح، قابل آزمون "
        "و قابل ردیابی را الزام می‌کند (International Organization for Standardization et al., 2018).",
    )

    add_h(doc, "۲.۲ سیستم‌های مشابه", 2)
    add_p(
        doc,
        "Moodle، Google Classroom، Canvas، openSIS و Fedena بررسی شدند. بسیاری بر محتوای "
        "درسی یا ثبت شاگرد تمرکز دارند، اما معاش محلی، مالیه دولتی و جریان مالی یکپارچه "
        "را به شکل مورد نیاز اکادمی‌های افغانستان پوشش نمی‌دهند (Costa et al., 2020).",
    )
    table_block(
        doc,
        "۲.۱",
        "مقایسه سیستم‌های موجود با NAMS",
        ["سیستم", "تمرکز", "محدودیت", "پاسخ NAMS"],
        [
            ["Moodle", "کورس و آزمون", "ضعیف در مالی محلی", "فیس، معاش، مالیه"],
            ["Google Classroom", "صنف آنلاین", "مدیریت کامل اکادمی ندارد", "ثبت، فیس، داشبورد"],
            ["Canvas", "LMS پیشرفته", "پیچیده برای اکادمی کوچک", "طراحی محلی ساده‌تر"],
            ["openSIS", "ثبت شاگرد", "معاش و مالیه ندارد", "Payroll و Tax"],
            ["NAMS", "عملیات یکپارچه", "AI پیش‌بینی آینده", "سیستم اختصاصی قابل توسعه"],
        ],
        [
            "جدول ۲.۱ نشان می‌دهد NAMS تکرار ساده پلتفرم‌های موجود نیست.",
            "ارزش آن در یکپارچه‌سازی آموزشی–مالی–اداری برای اکادمی خصوصی است.",
        ],
        [
            "این مقایسه توجیه ساخت سیستم اختصاصی را برای اکادمی نکته تقویت می‌کند.",
        ],
    )

    add_h(doc, "۲.۳ روندها و تکنالوژی‌ها", 2)
    add_p(
        doc,
        "روند حرکت از ابزارهای جدا به سمت پلتفرم یکپارچه با داشبورد و تحلیل داده "
        "مستند است (Siemens, 2013). Node.js، Express، MongoDB و React برای تیم کوچک "
        "و تکرار Agile مناسب‌اند (Bass et al., 2012; Sommerville, 2016).",
    )
    add_p(
        doc,
        "Agile Scrum با SDLC ساختاریافته ترکیب شد تا هم انعطاف بازخورد و هم "
        "مستندسازی آکادمیک حفظ شود (Dikert et al., 2016; Schwaber & Sutherland, 2020).",
    )

    add_h(doc, "۲.۴ خلاهای تحقیق", 2)
    add_p(
        doc,
        "خلاهای اصلی: (۱) پراکندگی جریان آموزشی و مالی؛ (۲) ضعف داده برای AI؛ "
        "(۳) نبود شفافیت مالی محلی؛ (۴) عدم تطبیق LMSهای بین‌المللی با مالیه و "
        "معاش افغانستان (Rahman & Saha, 2019). NAMS این خلاهای را در سه نمونه "
        "اولیه هدف قرار داده است.",
    )

    add_h(doc, "۲.۵ امنیت و پذیرش کاربر", 2)
    add_p(
        doc,
        "ISO/IEC 25010 امنیت را از خصوصیات کیفیت نرم‌افزار می‌داند "
        "(International Organization for Standardization & International Electrotechnical Commission, 2011). "
        "NAMS از JWT، RBAC و میان‌افزار صلاحیت استفاده می‌کند.",
    )
    add_p(
        doc,
        "مدل پذیرش فناوری Davis (1989) و Venkatesh et al. (2003) توضیح می‌دهد که "
        "مفید بودن درک‌شده و سهولت استفاده بر پذیرش اثر دارد؛ پرسشنامه ۲۰ اکادمی "
        "این متغیرها را سنجید.",
    )

    add_h(doc, "۲.۶ جمع‌بندی فصل", 2)
    add_p(
        doc,
        "ادبیات از یکپارچگی عملیاتی–مالی، داده ساختارمند برای analytics، Agile، "
        "معماری لایه‌ای و پذیرش کاربر حمایت می‌کند. فصل سوم نحوه عملیاتی‌سازی "
        "این اصول در NAMS را شرح می‌دهد.",
    )
    page_break(doc)


def build_chapter3(doc: Document):
    add_h(doc, "فصل سوم: میتودولوژی", 1)
    add_p(
        doc,
        "این فصل روش‌های تحقیق و توسعه طراحی و پیاده‌سازی NAMS را شرح می‌دهد: "
        "نیازسنجی، جامعه و نمونه، جمع‌آوری و تحلیل داده، نیازمندی‌های کارکردی و "
        "غیرکارکردی، Agile Scrum، سه نمونه اولیه، طراحی سیستم، پیاده‌سازی و آزمون "
        "(Creswell & Creswell, 2018; Pressman & Maxim, 2020).",
    )

    add_h(doc, "۳.۱ میتودولوژی تحقیق", 2)
    add_p(
        doc,
        "این مطالعه از طرح مختلط کاربردی (mixed-methods) استفاده کرد. داده کمی "
        "پرسشنامه تقاضای امکانات را سنجید؛ تحلیل کیفی جریان کار نقش‌ها را توضیح "
        "داد (DeLone & McLean, 2003).",
    )
    add_p(
        doc,
        "قبل از نهایی‌سازی پرسشنامه، فهرست اولیه امکانات از ادبیات فصل ۲ و "
        "بررسی Moodle، Google Classroom و Canvas استخراج شد (Aldowah et al., 2019; "
        "Romero & Ventura, 2020).",
    )

    add_h(doc, "۳.۱.۱ استخراج نیازمندی‌ها", 3)
    add_p(
        doc,
        "استخراج نیازمندی‌ها ترکیبی از (الف) بررسی پروپوزل، (ب) مشاهده جریان کار "
        "اکادمی، (ج) تحلیل نقش و صلاحیت، و (د) پرسشنامه ساختاریافته ۲۰ اکادمی بود "
        "(International Organization for Standardization et al., 2018).",
    )
    add_p(
        doc,
        "فهرست اولیه شامل ثبت شاگرد، صنف، تقسیم اوقات، حاضری دستی، امتحان، فیس، "
        "معاش، پیام و داشبورد بود. امکانات AI امتحان، لینک Google Meet و حاضری "
        "خودکار والدین به عنوان کاندید پس از نظرسنجی در نظر گرفته شد.",
    )
    add_p(
        doc,
        "دو لایه نیازمندی ایجاد شد: لایه عملیاتی پایه (Prototype 1) و لایه "
        "تقویتی پس از تأیید ۱۸ از ۲۰ اکادمی (Prototype 2) و مالیه دولتی "
        "درخواستی اکادمی نکته (Prototype 3).",
    )

    add_h(doc, "۳.۱.۲ جامعه و نمونه‌گیری", 3)
    table_block(
        doc,
        "۳.۱",
        "جامعه، نمونه و مشخصات پاسخ‌دهندگان NAMS",
        ["مورد", "توضیح", "تعداد / مقدار"],
        [
            ["جامعه هدف", "اکادمی‌های خصوصی با عملیات آموزشی–مالی", "بخش شهری"],
            ["اکادمی‌های دعوت‌شده", "تماس برای پرسشنامه", "۲۰"],
            ["پرسشنامه تکمیل‌شده", "پاسخ کامل", "۲۰"],
            ["درخواست امکانات Prototype 2", "AI امتحان، Meet، حاضری والدین", "۱۸"],
            ["اکادمی تازه‌تأسیس (بدون نظر)", "سابقه عملیاتی محدود", "۲"],
            ["پاسخ‌دهندگان اصلی", "اداره، مالک، استاد، مدیر شعبه", "۲۰ نفر"],
        ],
        [
            "جدول ۳.۱ چارچوب نمونه‌گیری را خلاصه می‌کند. ۱۸ اکادمی صریحاً بسته "
            "Prototype 2 را خواستند؛ ۲ تازه‌تأسیس نظر مشخص ندادند.",
            "نمونه کوچک اما غنی از نظر اطلاعات برای پروژه لیسانس است "
            "(Brooke, 1996; Creswell & Creswell, 2018).",
        ],
        [
            "تقسیم ۱۸ به ۲ قاعده تصمیم Product Owner برای Prototype 2 شد.",
        ],
    )

    add_h(doc, "۳.۱.۳ جمع‌آوری داده", 3)
    add_p(
        doc,
        "پرسشنامه ساختاریافته (ضمیمه الف) شامل ابزار فعلی، دردها، اولویت امکانات "
        "و احتمال پذیرش (Likert, 1932) بود. یادداشت‌های کیفی در جریان بررسی فیس، "
        "حاضری و معاش در اکادمی نکته جمع‌آوری شد.",
    )
    add_p(
        doc,
        "بخش‌های پرسشنامه به ماژول‌های طراحی نگاشت شدند تا ردیابی "
        "نیازمندی–منبع حفظ شود (International Organization for Standardization et al., 2018).",
    )

    add_h(doc, "۳.۱.۴ تحلیل و ارائه داده", 3)
    add_p(
        doc,
        "پاسخ‌ها با آمار توصیفی (فراوانی، درصد، میانگین اولویت) تحلیل شدند. "
        "نتایج در فصل چهارم با جداول و شکل‌های APA ارائه می‌شوند. "
        "یادداشت‌های کیفی به تم‌های نیازمندی کدگذاری شدند (Miles et al., 2014).",
    )

    add_h(doc, "۳.۱.۵ نیازمندی‌های کارکردی", 3)
    table_block(
        doc,
        "۳.۲",
        "نیازمندی‌های کارکردی منتخب NAMS",
        ["شناسه", "نیازمندی", "اولویت", "منبع", "آزمون"],
        [
            ["FR-01", "ثبت و مدیریت شاگرد با صنف، مضمون، استاد و فیس", "بالا", "جریان کار", "CRUD/API"],
            ["FR-02", "مدیریت استاد، نوع معاش و تخصیص", "بالا", "معاش", "Payroll"],
            ["FR-03", "ثبت حاضری جلسه‌ای و راپور", "بالا", "پرسشنامه", "Functional"],
            ["FR-04", "امتحان، نتیجه و انتشار برای شاگرد/والد", "بالا", "پروپوزل", "Exam"],
            ["FR-05", "فیس، مانده و ارتباط با معاش", "بالا", "مالی", "Payment"],
            ["FR-06", "JWT و RBAC", "بالا", "امنیت", "Auth"],
            ["FR-07", "رابط چندزبانه EN/FA/PS و RTL", "متوسط", "بافت محلی", "Locale"],
            ["FR-08", "تولید سوال AI (Prototype 2)", "متوسط", "۱۸/۲۰", "آزمون سرویس هوش مصنوعی"],
            ["FR-09", "لینک Meet و حاضری آنلاین خودکار (P2)", "متوسط", "۱۸/۲۰", "Job/policy"],
            ["FR-10", "گزارش مالیه فصل‌وار/سالانه (P3)", "بالا", "درخواست نکته", "Tax"],
        ],
        [
            "جدول ۳.۲ نیازمندی‌های قابل آزمون Product Backlog را تعریف می‌کند "
            "(Schwaber & Sutherland, 2020).",
            "FR-08 و FR-09 عمداً تا Prototype 2 به تعویق افتادند.",
        ],
        [
            "ستون آزمون پیوند narrative با شواهد هدفمند را تضمین می‌کند "
            "(Pressman & Maxim, 2020).",
        ],
    )

    add_h(doc, "۳.۱.۶ نیازمندی‌های غیرکارکردی", 3)
    table_block(
        doc,
        "۳.۳",
        "نیازمندی‌های غیرکارکردی منتخب NAMS",
        ["شناسه", "رده", "نیازمندی", "آزمون"],
        [
            ["NFR-01", "کارایی", "پاسخ CRUD و داشبورد در بار LAN عادی", "Timing"],
            ["NFR-02", "امنیت", "JWT، hash رمز، RBAC، CSRF", "Security"],
            ["NFR-03", "قابلیت استفاده", "منوی نقش‌محور، تم روشن/تاریک", "UAT"],
            ["NFR-04", "نگهداری", "Express modular و React feature-based", "Review"],
            ["NFR-05", "قابلیت اطمینان", "محاسبات مالی و مالیه قطعی", "Unit"],
            ["NFR-06", "سازگاری", "مرورگر دسکتاپ، Flutter/PWA", "Device"],
            ["NFR-07", "بومی‌سازی", "EN/FA/PS و فیلتر Hijri", "Locale"],
        ],
        [
            "جدول ۳.۳ کیفیت سیستمی را برای داده حساس شاگرد و مالی تعریف می‌کند "
            "(International Organization for Standardization & International Electrotechnical Commission, 2011).",
            "NFR در Sprint Review به عنوان چک‌لیست کیفیت استفاده شد.",
        ],
        [
            "امکانات نمونه‌های دوم و سوم بدون رعایت نیازمندی غیرکارکردی مربوط، تکمیل‌شده اعلام نمی‌شد.",
        ],
    )

    add_h(doc, "۳.۱.۷ نیازمندی‌های داده", 3)
    add_p(
        doc,
        "NAMS به رکوردهای ساختاریافته کاربر، نقش، شعبه، شاگرد، استاد، خانواده، "
        "صنف، مضمون، ثبت‌نام، تقسیم اوقات، حاضری، سیاست حاضری، امتحان، پرداخت، "
        "معاش، تنظیمات مالیه، پیام و audit نیاز دارد.",
    )
    add_p(
        doc,
        "MongoDB با مراجع ObjectId برای معاش، مالیه و گزارش‌دهی انتخاب شد "
        "(Chodorow, 2013; Connolly & Begg, 2015).",
    )

    add_h(doc, "۳.۲ میتودولوژی توسعه", 2)
    add_p(
        doc,
        "توسعه با Agile Scrum در پوشش SDLC انجام شد: تحلیل، طراحی، پیاده‌سازی، "
        "آزمون (Pressman & Maxim, 2020; Schwaber & Sutherland, 2020).",
    )

    add_h(doc, "۳.۲.۱ رویکرد Agile", 3)
    add_p(
        doc,
        "Agile چون اولویت‌ها پس از دمو روشن شدند انتخاب شد. Product Backlog، "
        "اسپرینت، بازبینی اسپرینت و تعریف انجام‌شدگی استفاده شد "
        "(Dikert et al., 2016).",
    )
    figure_block(
        doc,
        ASSETS / "fig_agile_scrum.png",
        "۳.۰",
        "چرخه Agile Scrum در سه نمونه NAMS",
        [
            "شکل ۳.۰ نگاشت چرخه Agile به تحویل نمونه‌ها را نشان می‌دهد. "
            "Backlog ابتدا عملیات پایه، سپس تقویت نظرسنجی، سپس مالیه دولتی را داشت.",
            "حذف عمدی AI و Meet در Prototype 1 تصمیم اولویت‌بندی بود، نه نقص.",
        ],
        [
            "پس از شکل ۳.۰، زیربخش‌های Scrum، تکرار و نمونه‌های اولیه می‌آیند.",
        ],
    )

    add_h(doc, "۳.۲.۲ سازمان Scrum", 3)
    add_p(
        doc,
        "نقش Product Owner بین تیم و مدیریت اکادمی نکته تقسیم شد. توسعه‌دهندگان "
        "بک‌اند، ویب و موبایل را پیاده کردند. آزمون واحد مالی و پذیرش کاربر در هر دروازه "
        "نمونه انجام شد (Schwaber & Sutherland, 2020).",
    )

    add_h(doc, "۳.۲.۳ تکرارهای توسعه", 3)
    add_p(
        doc,
        "ترتیب: (۱) احراز هویت و masters؛ (۲) حاضری و امتحان؛ (۳) پرداخت و معاش؛ "
        "(۴) پیام و داشبورد (P1 کامل)； (۵) AI، Meet، حاضری والدین (P2)； "
        "(۶) مالیه فصل‌وار/سالانه (P3) (Sommerville, 2016).",
    )

    add_h(doc, "۳.۲.۴ توسعه نمونه و بازبینی کاربر", 3)
    add_p(
        doc,
        "سه مرحله بازبینی‌شده: Prototype 1 خط پایه؛ Prototype 2 پس از ۱۸/۲۰؛ "
        "Prototype 3 مالیه دولتی درخواستی اکادمی نکته.",
    )

    add_h(doc, "۳.۲.۴.۱ نمونه اول — خط پایه یکپارچه", 3)
    add_p(
        doc,
        "Prototype 1 حداقل پلتفرم یکپارچه را تحویل داد: احراز هویت، RBAC، شاگرد، "
        "استاد، صنف، حاضری دستی، امتحان دستی، فیس، معاش پایه، پیام و داشبورد.",
    )
    add_p(
        doc,
        "عمداً خارج از Prototype 1: (الف) تولید سوال AI، (ب) Google Meet / "
        "لینک کلاس آنلاین به عنوان کانال حاضری، (ج) حاضری خودکار والدین "
        "(Pressman & Maxim, 2020).",
    )
    figure_block(
        doc,
        ASSETS / "fig_3_1_proto1_dashboard.png",
        "۳.۱",
        "داشبورد اداره Prototype 1",
        [
            "شکل ۳.۱ داشبورد عملیات پایه را نشان می‌دهد؛ بدون AI یا اتوماسیون Meet.",
            "بازبینان KPIهای ثبت‌نام، حاضری و فیس را در یک پرش تأیید کردند "
            "(Davis, 1989).",
        ],
        [
            "شکایت باقی‌مانده: سرعت امتحان و دید والدین — دقیقاً شکاف Prototype 2.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_2_proto1_students.png",
        "۳.۲",
        "مدیریت شاگرد، صنف، استاد و تقسیم اوقات Prototype 1",
        [
            "شکل ۳.۲ master-data آموزشی را نشان می‌دهد؛ پیوند شاگرد به صنف، مضمون، "
            "استاد و فیس برای محاسبه معاش پایدار است.",
            "بدون این پیوندها، ماژول مالی دوباره پراکندگی فصل ۱ را تکرار می‌کند "
            "(Connolly & Begg, 2015).",
        ],
        [
            "بازبینی شکل ۳.۲ قواعد اعتبارسنجی را قبل از P2 اصلاح کرد.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_3_proto1_attendance_exam.png",
        "۳.۳",
        "حاضری دستی و امتحان پایه Prototype 1",
        [
            "شکل ۳.۳ مسیر حاضری دستی و امتحان بدون AI را مستند می‌کند.",
            "اتوماسیون تا تأیید سیاست و نظرسنجی به تعویق افتاد.",
        ],
        [
            "این شکل خط پایه کنترل‌شده برای مقایسه با نمونه دوم در فصل چهارم است.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_4_proto1_finance.png",
        "۳.۴",
        "فیس و معاش پایه Prototype 1",
        [
            "شکل ۳.۴ فیس، مانده، مصارف و معاش پایه را نشان می‌دهد؛ "
            "بسته مالیه دولتی هنوز نیست.",
            "بازبینان مالی عملیات روزانه را پذیرفتند؛ compliance بعداً P3 شد.",
        ],
        [
            "شکل‌های ۳.۱–۳.۴ قرارداد Prototype 1 را تعریف می‌کنند.",
        ],
    )

    add_h(doc, "۳.۲.۴.۲ نمونه دوم — تقویت نظرسنجی", 3)
    add_p(
        doc,
        "پس از تحلیل پرسشنامه، ۱۸ از ۲۰ اکادمی AI امتحان، لینک Meet و "
        "حاضری خودکار والدین را خواستند. ۲ تازه‌تأسیس نظر مشخص ندادند.",
    )
    add_p(
        doc,
        "Prototype 2 افزود: (۱) تولید سوال AI با بازبینی استاد؛ (۲) لینک Meet "
        "در تقسیم اوقات؛ (۳) حاضری آنلاین خودکار و اعلان والدین بر اساس سیاست.",
    )
    add_p(
        doc,
        "امکانات پشت flag پیکربندی ماندند تا اکادمی‌های آماده‌نشده رفتار P1 "
        "را حفظ کنند (Dikert et al., 2016).",
    )
    figure_block(
        doc,
        ASSETS / "fig_3_5_proto2_ai_exam.png",
        "۳.۵",
        "تولید سوال امتحان با کمک AI — Prototype 2",
        [
            "شکل ۳.۵ workspace AI را نشان می‌دهد؛ استاد پیش‌نویس را ویرایش "
            "می‌کند (Romero & Ventura, 2020).",
            "مسئولیت علمی با استاد می‌ماند.",
        ],
        [
            "استادان صرفه‌جویی زمان در پیش‌نویس گزارش دادند.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_6_proto2_meet_attendance.png",
        "۳.۶",
        "لینک Meet و حاضری خودکار والدین — Prototype 2",
        [
            "شکل ۳.۶ پیوند تقسیم اوقات آنلاین با اتوماسیون حاضری را نشان می‌دهد.",
            "والدین حاضری را علامت نمی‌زنند؛ دید و هشدار دریافت می‌کنند.",
        ],
        [
            "شکاف ارتباطی فصل ۱ برای جلسات آنلاین بسته شد.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_7_proto2_survey_impact.png",
        "۳.۷",
        "توجیه Prototype 2 از پرسشنامه ۲۰ اکادمی",
        [
            "شکل ۳.۷ ۱۸ درخواست در برابر ۲ بدون نظر را خلاصه می‌کند.",
            "ردیابی نیازمندی از استخراج تا پیاده‌سازی حفظ می‌شود "
            "(International Organization for Standardization et al., 2018).",
        ],
        [
            "فصل ۴ همان داده را با نمودارهای کامل‌تر گسترش می‌دهد.",
        ],
    )

    add_h(doc, "۳.۲.۴.۳ نمونه سوم — گزارش مالیه دولتی", 3)
    add_p(
        doc,
        "Prototype 3 از درخواست خود اکادمی نکته (نه نظرسنجی چنداکادمی) "
        "شکل گرفت: گزارش مالیه فصل‌وار (Q1–Q4) و سالانه، BRT و مالیه معاش "
        "کارمندان، فیلتر هجری و صدور فایل PDF.",
    )
    add_p(
        doc,
        "معماری مالی modular اجازه داد بدون بازنویسی ماژول آموزشی اضافه شود "
        "(Bass et al., 2012).",
    )
    figure_block(
        doc,
        ASSETS / "fig_3_8_proto3_tax_dashboard.png",
        "۳.۸",
        "داشبورد مالیه دولتی Prototype 3",
        [
            "شکل ۳.۸ نمای فصل‌وار و سالانه را بدون Excel خارجی فراهم می‌کند.",
            "پاسخ مستقیم به درخواست sponsor است.",
        ],
        [
            "بازبینان تفکیک مالیه معاش و BRT را خواستند.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_9_proto3_quarterly.png",
        "۳.۹",
        "فیلتر و مجموع مالیه فصل‌وار Prototype 3",
        [
            "شکل ۳.۹ کنترل سال و فصل Hijri و مجموع فیس و BRT را نشان می‌دهد.",
            "ریتم فiling فصل‌وار اکادمی را پشتیبانی می‌کند.",
        ],
        [
            "مقایسه شعبه در یک فصل ممکن است.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_10_proto3_annual.png",
        "۳.۱۰",
        "بسته سالانه و صدور گزارش نمونه سوم",
        [
            "شکل ۳.۱۰ چهار فصل را در بسته سالانه ادغام می‌کند.",
            "خروجی چاپی برای بایگانی compliance است.",
        ],
        [
            "داستان نمونه‌ها: P1 عملیات، P2 تقاضای ۱۸/۲۰، P3 مالیه نکته.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_11_proto3_payroll_tax.png",
        "۳.۱۱",
        "یکپارچگی معاش–مالیه Prototype 3",
        [
            "شکل ۳.۱۱ رکورد معاش را به سرویس مالیه وصل می‌کند.",
            "شفافیت برای اعتماد institutional ضروری است (DeLone & McLean, 2003).",
        ],
        [
            "پس از شکل ۳.۱۱، طراحی رسمی در بخش ۳.۳ آغاز می‌شود.",
        ],
    )

    add_h(doc, "۳.۳ طراحی سیستم", 2)
    add_p(
        doc,
        "طراحی شامل use case، sequence، معماری و دیتابیس است "
        "(Fowler, 2004; Bass et al., 2012).",
    )

    add_h(doc, "۳.۳.۱ مدل Use Case", 3)
    figure_block(
        doc,
        ASSETS / "fig_3_12_use_case.png",
        "۳.۱۲",
        "نمودار Use Case سیستم مدیریت اکادمی نکته",
        [
            "شکل ۳.۱۲ actors اصلی: Student، Parent، Teacher، Admin، Super Admin، Owner.",
            "Use caseها شامل فیس، معاش، مالیه، Meet و AI در P2 هستند.",
            "مرز RBAC بخش محصول است، نه افزونه.",
        ],
        [
            "هر مورد کاربرد برای والدین باید محدوده API متناظر داشته باشد.",
            "زیربخش بعد sequence سه جریان پرریسک را جزئی می‌کند.",
        ],
    )

    add_h(doc, "۳.۳.۲ نمودارهای Sequence", 3)
    figure_block(
        doc,
        MONO / "Figure_3_13_Student_Enrollment_and_Fee_Payment_Sequence.png",
        "۳.۱۳",
        "Sequence ثبت‌نام شاگرد و پرداخت فیس",
        [
            "شکل ۳.۱۳ ثبت شاگرد، پیوند صنف/مضمون و پرداخت با اثر جانبی معاش را "
            "مدل می‌کند.",
            "پرداخت transaction ایزوله نیست؛ پایه tax و commission را تغذیه می‌کند.",
            "مرکز ادعای شفافیت مالی NAMS است.",
        ],
        [
            "این مسیر قبل از AI و tax در P1 موجود بود — ستون فقرات داده.",
        ],
    )
    figure_block(
        doc,
        MONO / "Figure_3_14_Attendance_and_Examination_Sequence.png",
        "۳.۱۴",
        "Sequence حاضری و امتحان",
        [
            "شکل ۳.۱۴ حاضری دستی P1 و اتوماسیون آنلاین P2 را در یک مدل دارد.",
            "جریان امتحان شامل AI اختیاری، انتشار و اعلان والدین است.",
        ],
        [
            "طراحی آزمون یکپارچه‌سازی برای مسیر نشر و اطلاع‌رسانی را پشتیبانی می‌کند.",
        ],
    )
    figure_block(
        doc,
        MONO / "Figure_3_15_Payroll_Tax_and_Financial_Report_Sequence.png",
        "۳.۱۵",
        "Sequence معاش، مالیه و گزارش مالی",
        [
            "شکل ۳.۱۵ محاسبه معاش، لایه‌های مالیه کارمند، مالیات دریافتی و صدور با فیلتر هجری "
            "را نشان می‌دهد.",
            "audit log تولیدکننده گزارش را نگه می‌دارد.",
        ],
        [
            "شکل‌های ۳.۱۳–۳.۱۵ سه مسیر پرریسک: پول ورودی، شواهد علمی، compliance خروجی.",
        ],
    )

    add_h(doc, "۳.۳.۳ طراحی معماری", 3)
    figure_block(
        doc,
        MONO / "Figure_3_16_NAMS_Architecture_Diagram.png",
        "۳.۱۶",
        "معماری لایه‌ای NAMS",
        [
            "شکل ۳.۱۶ React web، PWA/Flutter، Express API، MongoDB، notification، "
            "AI و reporting را نشان می‌دهد.",
            "JWT و RBAC لایه API را cross-cut می‌کنند (Bass et al., 2012).",
        ],
        [
            "jobs P2 و tax P3 بدون redesign client اضافه شدند.",
        ],
    )
    table_block(
        doc,
        "۳.۴",
        "تکنالوژی‌ها و توجیه NAMS",
        ["لایه", "تکنالوژی", "توجیه"],
        [
            ["Web", "React 18, Vite, TS, Tailwind", "SPA modular"],
            ["Mobile", "Flutter / PWA", "دسترسی تلفن"],
            ["API", "Node.js, Express, TS", "مسیرهای modular"],
            ["DB", "MongoDB, Mongoose", "سند با reference"],
            ["Auth", "JWT, bcrypt, RBAC", "دسترسی نقش‌محور"],
            ["Jobs", "node-cron", "اتوماسیون P2"],
            ["AI", "OpenAI/Groq/Gemini", "امتحان اختیاری"],
            ["Report", "PDF/print", "عملیاتی و tax"],
        ],
        [
            "جدول ۳.۴ انتخاب پشته تکنالوژی را با شکل ۳.۱۶ هم‌تراز می‌کند.",
            "P1 بدون کلید AI operable می‌ماند.",
        ],
        [
            "فصل پنجم پشته تکنالوژی را با سیستم‌های مدیریت یادگیری فاقد معاش و مالیه مقایسه می‌کند.",
        ],
    )

    add_h(doc, "۳.۳.۴ طراحی دیتابیس", 3)
    figure_block(
        doc,
        MONO / "Figure_3_17_NAMS_Entity_Relationship_Diagram.png",
        "۳.۱۷",
        "نمای ERD مجموعه‌های اصلی NAMS",
        [
            "شکل ۳.۱۷ users، students، payments، salary و messages را با reference "
            "نشان می‌دهد (Chodorow, 2013).",
            "soft-delete و timestamp برای audit فصل‌وار.",
        ],
        [
            "با شکل ۳.۱۸ schema سند تکمیل می‌شود.",
        ],
    )
    figure_block(
        doc,
        MONO / "Figure_3_18_NAMS_Core_Relational_Document_Schema.png",
        "۳.۱۸",
        "Schema سند–رابطه‌ای هسته NAMS",
        [
            "شکل ۳.۱۸ فیلدهای کلیدی و onlineAutoMarkEnabled برای P2 را نشان می‌دهد.",
            "tax از operational collections می‌خواند، ledger موازی ندارد.",
        ],
        [
            "پس از طراحی، بخش ۳.۴ پیاده‌سازی را خلاصه می‌کند.",
        ],
    )

    add_h(doc, "۳.۴ پیاده‌سازی", 2)
    add_p(
        doc,
        "Backend بر اساس domain route/service؛ frontend feature-based؛ mobile "
        "همان REST API. CI lint/build در هر sprint.",
    )

    add_h(doc, "۳.۴.۱ محیط توسعه", 3)
    add_p(
        doc,
        "Node.js 18+، MongoDB محلی، React/Vite، TypeScript، Git و .env برای "
        "secrets. اسکریپت offline Windows برای demo LAN.",
    )

    add_h(doc, "۳.۴.۲ فازهای پیاده‌سازی", 3)
    add_p(
        doc,
        "فاز A: هویت و masters؛ B: حاضری/امتحان؛ C: پرداخت/معاش؛ D: پیام/"
        "داشبورد (نمونه ۱)؛ مرحله بعد: هوش مصنوعی، جلسه آنلاین و حاضری والدین (نمونه ۲)؛ سپس مالیه (نمونه ۳). "
        "هر فاز با checklist FR/NFR بسته شد (Schwaber & Sutherland, 2020).",
    )

    add_h(doc, "۳.۴.۳ ماژول‌های کلیدی", 3)
    add_p(
        doc,
        "auth، masters، enrollment، attendance policy، exams/AI، payments، "
        "معاش و مالیه، پیام‌ها، گزارش‌ها، سجل حسابرسی و داشبوردها.",
    )

    add_h(doc, "۳.۴.۴ چالش‌ها و راه‌حل‌ها", 3)
    table_block(
        doc,
        "۳.۵",
        "چالش‌های پیاده‌سازی NAMS",
        ["چالش", "اثر", "راه‌حل"],
        [
            ["تغییر اولویت پس از دمو", "ریسک محدوده", "Backlog + 3 gate نمونه"],
            ["معاش فیصدی وابسته به پرداخت", "اختلاف محاسبه", "سرویس deterministic + unit"],
            ["RTL چندزبانه", "خرابی layout", "locale + تست RTL"],
            ["اعتماد حاضری آنلاین", "شکایت والد", "policy flag + audit"],
            ["وضوح قواعد tax", "ریسک compliance", "کارگاه sponsor + P3"],
            ["کیفیت سوال AI", "بی‌اعتمادی استاد", "پیش‌نویس قابل ویرایش"],
        ],
        [
            "جدول ۳.۵ بیشتر مسائل سازمانی–نیازمندی بود تا blocker فنی "
            "(Dikert et al., 2016).",
            "tax از workshop sponsor قابل حدس نبود.",
        ],
        [
            "ستون راه‌حل به استراتژی آزمون بخش ۳.۵ وصل است.",
        ],
    )

    add_h(doc, "۳.۵ آزمون", 2)
    add_p(
        doc,
        "آزمون واحد (مالی و مالیه)، یکپارچه‌سازی (ثبت–پرداخت–معاش، امتحان–اطلاع) "
        "و پذیرش کاربر در هر دروازه نمونه (Sommerville, 2016).",
    )

    add_h(doc, "۳.۵.۱ آزمون واحد", 3)
    table_block(
        doc,
        "۳.۶",
        "نمونه Test Case واحد NAMS",
        ["شناسه", "حوزه", "سناریو", "نتیجه مورد انتظار"],
        [
            ["UT-01", "Payroll", "معاش فیصدی روی پرداخت شاگرد", "Gross = base × %"],
            ["UT-02", "Tax", "معاش در bracket", "Tax مطابق قاعده"],
            ["UT-03", "BRT", "مجموع فیس فصل", "BRT = % tuition"],
            ["UT-04", "Policy", "onlineAutoMark=false", "بدون mark خودکار"],
            ["UT-05", "Policy", "onlineAutoMark=true", "mark + notify والد"],
            ["UT-06", "AI", "MCQ تولید و ویرایش", "publish متن نهایی"],
            ["UT-07", "RBAC", "توکن شاگرد به گزارش معاش", "رد"],
            ["UT-08", "Validation", "ثبت‌نام تکراری", "خطای واضح"],
        ],
        [
            "جدول ۳.۶ UT-04/05 رفتار P1 در برابر P2 را قفل می‌کند.",
            "UT-02/03 از regression tax محافظت می‌کند.",
        ],
        [
            "قبل از بازبینی اسپرینت و اعلام تکمیل اجرا شد.",
        ],
    )

    add_h(doc, "۳.۵.۲ آزمون یکپارچه‌سازی", 3)
    add_p(
        doc,
        "مسیرهای چندماژولی: ثبت تا پرداخت؛ حاضری تا راپور؛ publish نتیجه؛ "
        "نشست آنلاین تا اطلاع‌رسانی؛ معاش تا صدور گزارش مالیه.",
    )

    add_h(doc, "۳.۵.۳ آزمون پذیرش کاربر", 3)
    add_p(
        doc,
        "پذیرش نمونه ۱: عملیات روزانه بدون اکسل. پذیرش نمونه ۲: هوش مصنوعی، جلسه آنلاین، حاضری والدین. "
        "پذیرش نمونه ۳: خروجی مالیه مطابق انتظار مدیریت. "
        "الهام از Brooke (1996) و Nielsen (1994).",
    )
    add_p(
        doc,
        "P1 پذیرفته شد؛ P2 برای ۱۸/۲۰ جذاب؛ P3 tax برای workflow compliance "
        "نکته پذیرفته شد.",
    )
    page_break(doc)


def build_chapter4(doc: Document):
    add_h(doc, "فصل چهارم: نتایج", 1)
    add_p(
        doc,
        "این فصل نتایج پرسشنامه ۲۰ اکادمی و دستاوردهای مهندسی سه نمونه اولیه "
        "را ارائه می‌کند — پروفایل پاسخ‌دهنده، ابزار فعلی، اولویت امکانات، "
        "اثر درک‌شده و خلاصه.",
    )

    add_h(doc, "۴.۱ نتایج پرسشنامه", 2)
    add_h(doc, "۴.۱.۱ پروفایل پاسخ‌دهندگان", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_1_respondent_role.png",
        "۴.۱",
        "نقش پاسخ‌دهندگان در ۲۰ اکادمی",
        [
            "شکل ۴.۱ نشان می‌دهد بیشتر پاسخ‌ها از اداره و مالک بود — مناسب "
            "برای تصمیم پذیرش ابزار و جریان مالی.",
            "استاد و مدیر شعبه دید عملیاتی حاضری و امتحان را آوردند.",
            "تنوع نقش از one-sided wishlist جلوگیری می‌کند (Creswell & Creswell, 2018).",
        ],
        [
            "پروفایل شکل ۴.۱ تفسیر اولویت‌ها را institutional می‌کند.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_4_2_academy_maturity.png",
        "۴.۲",
        "بلوغ عملیاتی اکادمی‌های شرکت‌کننده",
        [
            "شکل ۴.۲ ۱۸ اکادمی بیش از یک سال را از ۲ تازه‌تأسیس جدا می‌کند.",
            "۲ تازه‌تأسیس نظر مشخص درباره Prototype 2 ندادند.",
            "اکادمی بالغ درد Excel را concrete ارزیابی کرد.",
        ],
        [
            "قاعده ۱۸/۲۰ به جای unanimity توجیه می‌شود.",
        ],
    )

    add_h(doc, "۴.۱.۲ ابزار و روش فعلی", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_3_current_tools.png",
        "۴.۳",
        "ابزار آموزشی–اداری فعلی اکادمی‌ها",
        [
            "شکل ۴.۳ بیشتر وابستگی به کاغذ/Excel یا messaging را نشان می‌دهد.",
            "chat گروهی audit trail و tax packaging ندارد.",
            "توجیه MIS یکپارچه به جای app تک‌منظوره.",
        ],
        [
            "onboarding باید عادت Excel را تدریجی جایگزین کند — Agile phased.",
        ],
    )

    add_h(doc, "۴.۱.۳ علاقه و اولویت امکانات", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_4_feature_need.png",
        "۴.۴",
        "اکادمی‌های درخواست‌کننده بسته Prototype 2",
        [
            "شکل ۴.۴ نتیجه محوری P2: ۱۸ نیاز، ۲ بدون نظر — ۹۰٪ نمونه.",
            "Authority Product Owner برای scheduling P2 پس از P1.",
            "به workload استاد و اعتماد والدین وصل است.",
        ],
        [
            "با شکل ۴.۵ (میانگین اولویت گروه امکانات) بخوانید.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_4_5_priority.png",
        "۴.۵",
        "میانگین اولویت گروه‌های امکانات NAMS",
        [
            "شکل ۴.۵ فیس/معاش و حاضری در صدر؛ exams/AI و portal والدین نزدیک.",
            "tax بالا حتی قبل از کد P3 — تقاضای latent.",
            "suite یکپارچه خواسته می‌شود نه silver bullet (DeLone & McLean, 2003).",
        ],
        [
            "اولویت‌ها ترتیب sprint تحت فشار تقویم را هدایت کرد.",
        ],
    )

    add_h(doc, "۴.۱.۴ اثر درک‌شده", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_6_impact.png",
        "۴.۶",
        "توافق با گزاره‌های اثر مورد انتظار",
        [
            "شکل ۴.۶: تقریباً همه کاهش کار دستی و شفافیت را باور دارند.",
            "آمادگی tax قبل از P3 بالا — sponsor بعداً explicit کرد.",
            "perceived usefulness پیش‌بین پذیرش است (Davis, 1989; Venkatesh et al., 2003).",
        ],
        [
            "انتظارات است؛ proof longitudinal در فصل ۵ تمایز داده می‌شود.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_4_7_acceptance.png",
        "۴.۷",
        "احتمال پذیرش NAMS",
        [
            "شکل ۴.۷: ۱۲ very likely، ۶ likely، ۲ neutral؛ هیچ unlikely.",
            "neutral با ۲ تازه‌تأسیس و بدون نظر P2 هم‌پوشانی دارد.",
            "پذیرش قوی برای نمونه هدفمند.",
        ],
        [
            "پایان زیربخش پرسشنامه؛ نتایج مهندسی در خلاصه ۴.۱.۵.",
        ],
    )

    add_h(doc, "۴.۱.۵ خلاصه یافته‌ها", 3)
    add_p(
        doc,
        "یافته‌های پرسشنامه: ابزار پراکنده، تقاضای قوی فیس/حاضری/امتحان، "
        "۱۸/۲۰ برای P2، اثر بالا و پذیرش بالا. یافته‌های مهندسی: P1 usable، "
        "نمونه ۲ با پرچم پیکربندی، نمونه ۳ مالیه در پذیرش حامی، آزمون‌های واحد معاش و مالیه، "
        "diagrams هم‌تراز با monorepo.",
    )
    table_block(
        doc,
        "۴.۱",
        "نگاشت شواهد نتیجه به اهداف پروژه",
        ["تم هدف", "شواهد", "در gate"],
        [
            ["متمرکزسازی", "CRUD + ERD/schema", "P1"],
            ["دید حاضری", "دستی P1؛ auto P2", "P1–P2"],
            ["سرعت امتحان", "AI + بازبینی استاد", "P2"],
            ["کلاس آنلاین", "Meet + jobs", "P2"],
            ["شفافیت مالی", "پرداخت–معاش", "P1"],
            ["مالیه دولتی", "بسته فصل/سال", "P3"],
            ["پذیرش", "۱۸/۲۰ demand؛ ۱۸/۲۰ likely+", "Survey"],
        ],
        [
            "جدول ۴.۱ synthesis برای گذار به فصل ۵ است.",
            "هر ردیف به artifact یا survey اشاره دارد.",
        ],
        [
            "فصل پنجم از این جدول به‌عنوان فهرست راهنما استفاده می‌کند.",
        ],
    )
    page_break(doc)


def build_chapter5(doc: Document):
    add_h(doc, "فصل پنجم: بحث", 1)
    add_p(
        doc,
        "این فصل نتایج را تفسیر می‌کند: یافته‌های اصلی، ارتباط با اهداف و "
        "سوالات، ادبیات، پیامدهای عملی و محدودیت‌ها.",
    )

    add_h(doc, "۵.۱ یافته‌های اصلی", 2)
    add_p(
        doc,
        "سه یافته غالب: (۱) ابزار پراکنده — MIS یکپارچه gap واقعی است؛ "
        "(۲) تقاضا گسترده اما زمان‌بندی‌شده — P1 فوری، P2 پس از ۱۸/۲۰؛ "
        "(۳) tax (P3) capability حیاتی sponsor که LMS عمومی ندارد "
        "(Rahman & Saha, 2019).",
    )
    add_p(
        doc,
        "استراتژی تکراری از big-bang AI+tax در demo اول جلوگیری کرد "
        "(Dikert et al., 2016; Pressman & Maxim, 2020).",
    )

    add_h(doc, "۵.۲ ارتباط با اهداف و سوالات", 2)
    add_p(
        doc,
        "اهداف فصل ۱: متمرکزسازی (P1)، ارتباط/نظارت (P2)، compliance (P3)، "
        "analytics readiness جزئی با داده ساختارمند (Aldowah et al., 2019). "
        "پاسخ در سطح سیستم engineered + stakeholder evidence است "
        "(Hevner et al., 2004).",
    )
    table_block(
        doc,
        "۵.۱",
        "ماتریس پاسخ سوالات تحقیق",
        ["تم سوال", "پاسخ از شواهد NAMS"],
        [
            ["سوابق پراکنده", "مدل یکپارچه student–fee–attendance–exam"],
            ["معلومات به‌موقع", "داشبورد و portal نقش‌محور"],
            ["ارتباط", "پیام + alert حاضری P2"],
            ["analytics", "تاریخچه ساختارمند؛ ML آینده"],
            ["اتوماسیون", "پیوند معاش؛ کارهای پس‌زمینه نمونه‌های ۲ و ۳"],
        ],
        [
            "جدول ۵.۱ ادعاها را برای دفاع فشرده می‌کند.",
            "هر پاسخ در محدوده نمونه آماری و محدوده نمونه‌های سیستم است.",
        ],
        [
            "mixed-methods + Agile دلیل نوع شواهد هر پاسخ است.",
        ],
    )

    add_h(doc, "۵.۳ ادبیات و سیستم‌های موجود", 2)
    add_p(
        doc,
        "NAMS بیشتر پلتفرم عملیات اکادمی است تا LMS محتوا. LMSها اغلب finance "
        "خارجی می‌خواهند (Costa et al., 2020). NAMS عملیات را مرکز و "
        "learning-support را ضمیمه می‌کند.",
    )
    add_p(
        doc,
        "موفقیت سیستم معلوماتی: دقت معاش و مالیه اعتماد می‌سازد (DeLone & McLean, 2003). "
        "TAM: کنترل استاد بر AI برای trust لازم است (Davis, 1989).",
    )

    add_h(doc, "۵.۴ پیامدهای عملی", 2)
    add_p(
        doc,
        "مسیر Excel/chat به سیستم نقش‌آگاه با رشد تدریجی: P1 سپس flag P2. "
        "برای نکته P3 spreadsheet tax را کاهش می‌دهد. برای ۱۸/۲۰، بسته "
        "بسته هوش مصنوعی، جلسه آنلاین و حاضری والدین قابل انتقال به اکادمی‌های دیگر است.",
    )

    add_h(doc, "۵.۵ محدودیت‌ها", 2)
    add_p(
        doc,
        "محدودیت‌ها: نمونه هدفمند ۲۰، انتظارات self-report نه longitudinal، "
        "tax بر تفسیر operational sponsor، AI وابسته provider و انضباط استاد. "
        "احراز هویت زیستی و یکپارچه‌سازی ملی خارج از محدوده است.",
    )
    add_p(
        doc,
        "این محدودیت‌ها claim را برای پروژه کاربردی لیسانس bounded نگه "
        "می‌دارند (Hevner et al., 2004; Creswell & Creswell, 2018).",
    )
    page_break(doc)


def build_chapter6(doc: Document):
    add_h(doc, "فصل ششم: نتیجه‌گیری و کار آینده", 1)

    add_h(doc, "۶.۱ خلاصه کار انجام‌شده", 2)
    add_p(
        doc,
        "NAMS به عنوان پلتفرم یکپارچه ویب–API–موبایل برای عملیات اکادمی خصوصی "
        "طراحی و پیاده‌سازی شد. فصل ۱–۲ مسئله و ادبیات؛ فصل ۳ mixed-methods "
        "و Agile با سه نمونه؛ فصل ۴–۵ نتایج و تفسیر.",
    )
    add_p(
        doc,
        "نمونه ۱ بدون امتحان هوش مصنوعی، بدون اتوماسیون حاضری جلسه آنلاین، بدون مالیه کامل. "
        "P2 پس از ۱۸/۲۰. P3 مالیه فصل‌وار/سالانه درخواست نکته.",
    )

    add_h(doc, "۶.۲ مشارکت‌های کلیدی", 2)
    add_p(
        doc,
        "مشارکت‌ها: (۱) MIS یکپارچه آموزشی–مالی؛ (۲) استراتژی نمونه مبتنی "
        "بر ۲۰ اکادمی؛ (۳) حاضری آنلاین والدین با policy؛ (۴) AI exam با "
        "تأیید انسان؛ (۵) بسته tax از داده عملیاتی.",
    )

    add_h(doc, "۶.۳ دستیابی به اهداف", 2)
    table_block(
        doc,
        "۶.۱",
        "وضعیت دستیابی به اهداف NAMS",
        ["تم هدف", "وضعیت", "شواهد"],
        [
            ["مدیریت متمرکز", "تحقق یافته", "ماژول‌ها + مدل داده"],
            ["ارتباط", "تحقق (تقویت P2)", "پیام + alert والد"],
            ["نظارت علمی", "تحقق", "portal نتایج/حاضری"],
            ["کارایی عملیاتی", "تحقق", "P1 + automation P2"],
            ["آمادگی AI/analytics", "جزئی", "داده + AI authoring"],
            ["گزارش مالیه", "تحقق‌یافته (نمونه ۳)", "پذیرش مالیه"],
        ],
        [
            "«جزئی» برای ML predictor از over-claim جلوگیری می‌کند (Peffers et al., 2007).",
            "tax برای بسته sponsor-marked achieved است.",
        ],
        [
            "جدول ۶.۱ نشان می‌دهد اهداف در محدوده تعیین‌شده محقق شدند.",
        ],
    )

    add_h(doc, "۶.۴ پیشنهادات آینده", 2)
    add_p(
        doc,
        "آینده: pilot longitudinal، مدل at-risk، sync offline عمیق‌تر، parity "
        "Flutter، backup drill، accessibility audit، tax pack قابل پیکربندی "
        "برای jurisdiction، scoring کیفیت سوال AI، نمونه چندشهری بزرگ‌تر "
        "(Creswell & Creswell, 2018).",
    )
    page_break(doc)


def build_references(doc: Document):
    add_h(doc, "منابع", 1)
    refs = [
        "Aldowah, H., Al-Samarraie, H., & Fauzy, W. M. (2019). Educational data mining and learning analytics for 21st century higher education: A review and synthesis. Telematics and Informatics, 37, 13–49. https://doi.org/10.1016/j.tele.2019.01.007",
        "American Psychological Association. (2020). Publication manual of the American Psychological Association (7th ed.). American Psychological Association.",
        "Bass, L., Clements, P., & Kazman, R. (2012). Software architecture in practice (3rd ed.). Addison-Wesley.",
        "Brooke, J. (1996). SUS: A quick and dirty usability scale. In P. W. Jordan, B. Thomas, B. A. Weerdmeester, & I. L. McClelland (Eds.), Usability evaluation in industry (pp. 189–194). Taylor & Francis.",
        "Chodorow, K. (2013). MongoDB: The definitive guide (2nd ed.). O'Reilly Media.",
        "Connolly, T. M., & Begg, C. E. (2015). Database systems: A practical approach to design, implementation, and management (6th ed.). Pearson.",
        "Costa, C., Teixeira, L., & Alvelos, H. (2020). Exploring the usage of learning management systems in higher education: A systematic review. Education and Information Technologies, 25, 1–20.",
        "Creswell, J. W., & Creswell, J. D. (2018). Research design: Qualitative, quantitative, and mixed methods approaches (5th ed.). SAGE Publications.",
        "Davis, F. D. (1989). Perceived usefulness, perceived ease of use, and user acceptance of information technology. MIS Quarterly, 13(3), 319–340. https://doi.org/10.2307/249008",
        "DeLone, W. H., & McLean, E. R. (2003). The DeLone and McLean model of information systems success: A ten-year update. Journal of Management Information Systems, 19(4), 9–30. https://doi.org/10.1080/07421222.2003.11045748",
        "Dikert, K., Paasivaara, M., & Lassenius, C. (2016). Challenges and success factors for large-scale agile transformations: A systematic literature review. Journal of Systems and Software, 119, 87–108. https://doi.org/10.1016/j.jss.2016.06.013",
        "Fowler, M. (2004). UML distilled: A brief guide to the standard object modeling language (3rd ed.). Addison-Wesley.",
        "Hevner, A. R., March, S. T., Park, J., & Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75–105.",
        "International Organization for Standardization, & International Electrotechnical Commission. (2011). Systems and software engineering — Systems and software Quality Requirements and Evaluation (SQuaRE) — System and software quality models (ISO/IEC 25010:2011).",
        "International Organization for Standardization, International Electrotechnical Commission, & Institute of Electrical and Electronics Engineers. (2018). Systems and software engineering — Life cycle processes — Requirements engineering (ISO/IEC/IEEE 29148:2018).",
        "Likert, R. (1932). A technique for the measurement of attitudes. Archives of Psychology, 22(140), 1–55.",
        "Miles, M. B., Huberman, A. M., & Saldaña, J. (2014). Qualitative data analysis: A methods sourcebook (3rd ed.). SAGE Publications.",
        "Nielsen, J. (1994). Usability engineering. Morgan Kaufmann.",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45–77.",
        "Pressman, R. S., & Maxim, B. R. (2020). Software engineering: A practitioner's approach (9th ed.). McGraw-Hill Education.",
        "Rahman, M. M., & Saha, S. (2019). An integrated education management information system for private academic institutions. International Journal of Information Engineering and Electronic Business, 11(2), 1–10.",
        "Romero, C., & Ventura, S. (2020). Educational data mining and learning analytics: An updated survey. WIREs Data Mining and Knowledge Discovery, 10(3), e1355. https://doi.org/10.1002/widm.1355",
        "Schwaber, K., & Sutherland, J. (2020). The Scrum guide: The definitive guide to Scrum: The rules of the game. Scrum.org.",
        "Siemens, G. (2013). Learning analytics: The emergence of a discipline. American Behavioral Scientist, 57(10), 1380–1400. https://doi.org/10.1177/0002764213498851",
        "Sommerville, I. (2016). Software engineering (10th ed.). Pearson.",
        "Venkatesh, V., Morris, M. G., Davis, G. B., & Davis, F. D. (2003). User acceptance of information technology: Toward a unified view. MIS Quarterly, 27(3), 425–478.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        run = p.add_run(ref)
        set_run(run, size=BODY_PT, font=FONT_FALLBACK)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.first_line_indent = Inches(-0.5)
        pf.left_indent = Inches(0.5)
        pf.space_after = Pt(8)


def build_appendices(doc: Document):
    page_break(doc)
    add_h(doc, "ضمیمه الف: پرسشنامه‌ها", 1)
    add_h(doc, "الف.۱ پرسشنامه استخراج نیازمندی‌ها", 2)
    add_p(
        doc,
        "پرسشنامه به یک پاسخ‌دهنده اصلی در هر اکادمی (n=20) داده شد. "
        "Likert پنج‌درجه‌ای (Likert, 1932). خلاصه بخش‌ها:",
        first_indent=True,
    )
    add_p(doc, "بخش الف — پروفایل اکادمی: سن institution، تعداد تقریبی شاگرد، نقش پاسخ‌دهنده.", first_indent=True)
    add_p(doc, "بخش ب — ابزار فعلی: کاغذ/Excel، messaging، نرم‌افزار جزئی، هیچ.", first_indent=True)
    add_p(doc, "بخش ج — دردهای عملیاتی: ثبت‌نام، حاضری، امتحان، فیس، والدین، راپور، tax.", first_indent=True)
    add_p(doc, "بخش د — مفید بودن و اولویت ماژول‌های MIS (۱–۵).", first_indent=True)
    add_p(
        doc,
        "بخش ه — بسته Prototype 2: نیاز به AI امتحان، لینک Meet، حاضری خودکار والدین "
        "(بله / نظر مشخص ندارم).",
        first_indent=True,
    )
    add_p(doc, "بخش و — گزاره اثر و احتمال پذیرش.", first_indent=True)
    add_p(
        doc,
        "۱۸ اکادمی «بله»؛ ۲ تازه‌تأسیس «نظر مشخص ندارم» — محدوده نمونه دوم را "
        "مجاز کرد.",
        first_indent=True,
    )

    page_break(doc)
    add_h(doc, "ضمیمه ب: مخزن کد منبع", 1)
    add_p(
        doc,
        "کد NAMS monorepo: backend (Node/Express/TS)، frontend (React/Vite)، "
        "mobile/PWA، ERD، scripts تولید مونوگراف. MongoDB و .env محلی.",
        first_indent=True,
    )
    add_p(
        doc,
        "مسیرهای کلیدی: backend/src، frontend/src/features، ERD/، scripts/.",
        first_indent=True,
    )
    add_p(
        doc,
        "اعتبار demo از کانال تحویل دیپارتمنت جداگانه — در متن عمومی چاپ نمی‌شود.",
        first_indent=True,
    )

    page_break(doc)
    add_h(doc, "ضمیمه ج: داده خام آزمون", 1)
    add_p(
        doc,
        "خروجی CI و log آزمون در repository. این ضمیمه خلاصه جدول ۳.۶ را "
        "تأیید می‌کند.",
        first_indent=True,
    )
    table_block(
        doc,
        "ج.۱",
        "خلاصه نتیجه خام آزمون نماینده",
        ["Suite", "تمرکز", "نتیجه در gate P3"],
        [
            ["Finance unit", "Payroll % + tax bracket", "Pass"],
            ["Academy tax", "مجموع BRT فصل", "Pass"],
            ["Attendance policy", "auto-mark on/off", "Pass"],
            ["Auth/RBAC", "رد نقش", "Pass"],
            ["Exam/AI smoke", "generate-edit-publish", "Pass"],
        ],
        [
            "جدول ج.۱ digest فشرده برای examiner است.",
            "Pass = suite موفق در commit ارزیابی P3.",
        ],
        [
            "سجل کامل از یکپارچه‌سازی مداوم قابل صدور است بدون تغییر متن فصول.",
        ],
    )

    page_break(doc)
    add_h(doc, "ضمیمه د: راهنمای کاربر", 1)
    add_p(
        doc,
        "راهنمای مختصر برای جهت‌دهی examiner و onboarding — جایگزین help درون‌برنامه‌ای نیست.",
        first_indent=True,
    )
    add_h(doc, "د.۱ اداره", 2)
    add_p(
        doc,
        "۱) ورود admin. ۲) شعبه/صنف/مضمون. ۳) ثبت شاگرد و استاد. ۴) فیس و "
        "داشبورد. ۵) نمونه ۳: معاش و مالیه، فصل یا سال هجری، صدور PDF.",
        first_indent=True,
    )
    add_h(doc, "د.۲ استاد", 2)
    add_p(
        doc,
        "۱) صنف‌های assigned. ۲) حاضری. ۳) امتحان؛ AI draft (P2) و ویرایش. "
        "۴) نتایج. ۵) لینک جلسه آنلاین ترکیبی.",
        first_indent=True,
    )
    add_h(doc, "د.۳ شاگرد", 2)
    add_p(doc, "تقسیم اوقات، حاضری، نتایج، فیس، اعلان.", first_indent=True)
    add_h(doc, "د.۴ والدین", 2)
    add_p(
        doc,
        "portal خانواده؛ حاضری و نتایج فرزند؛ P2: notify حاضری آنلاین.",
        first_indent=True,
    )
    add_h(doc, "د.۵ مالک / سوپر ادمین", 2)
    add_p(
        doc,
        "داشبورد اکادمی؛ نقش و صلاحیت؛ صدور و بایگانی گزارش مالیه فصل‌وار و سالانه.",
        first_indent=True,
    )
    add_p(doc, "پایان ضمیمه‌های مونوگراف.", first_indent=False, align="center")


def main():
    if OUT.exists():
        OUT.unlink()
    doc = Document()
    configure_styles(doc)

    required = [
        ASSETS / "fig_agile_scrum.png",
        ASSETS / "fig_3_1_proto1_dashboard.png",
        ASSETS / "fig_3_12_use_case.png",
        ASSETS / "fig_4_1_respondent_role.png",
        MONO / "Figure_3_13_Student_Enrollment_and_Fee_Payment_Sequence.png",
        MONO / "Figure_3_16_NAMS_Architecture_Diagram.png",
        MONO / "Figure_3_17_NAMS_Entity_Relationship_Diagram.png",
        MONO / "Figure_3_18_NAMS_Core_Relational_Document_Schema.png",
    ]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing assets:\n" + "\n".join(missing))

    build_front_matter(doc)
    build_chapter1(doc)
    build_chapter2(doc)
    build_chapter3(doc)
    build_chapter4(doc)
    build_chapter5(doc)
    build_chapter6(doc)
    build_references(doc)
    build_appendices(doc)

    doc.save(str(OUT))
    size_mb = OUT.stat().st_size / (1024 * 1024)
    print("CREATED", str(OUT).encode("unicode_escape").decode())
    print("SIZE_MB", round(size_mb, 2))


if __name__ == "__main__":
    main()

