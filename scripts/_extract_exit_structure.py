# -*- coding: utf-8 -*-
from pathlib import Path
import re
import fitz
from docx import Document

root = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
mono = None
for p in root.iterdir():
    if p.is_dir() and (p / "exit.pdf").exists():
        mono = p
        break
if not mono:
    for p in root.rglob("exit.pdf"):
        mono = p.parent
        break
print("MONO_OK", mono is not None)
print("MONO_PATH", str(mono).encode("unicode_escape").decode())

doc = fitz.open(str(mono / "exit.pdf"))
print("pages", doc.page_count)

text_parts = []
for i in range(min(50, doc.page_count)):
    t = doc.load_page(i).get_text("text")
    text_parts.append(f"\n===== PAGE {i+1} =====\n" + t)
(mono / "_exit_toc_extract.txt").write_text("\n".join(text_parts), encoding="utf-8")

heads = []
for i in range(doc.page_count):
    t = doc.load_page(i).get_text("text")
    for line in t.splitlines():
        s = line.strip()
        if not s or len(s) > 140:
            continue
        if re.match(
            r"^(Chapter|CHAPTER|Appendix|APPENDIX|References|REFERENCES|Bibliography|"
            r"List of|TABLE OF|Table of|\d+\.\d+(\.\d+)?\s|[3-9]\.\s|CHAPTER\s)",
            s,
        ):
            heads.append(f"p{i+1}: {s}")

uniq = []
seen = set()
for h in heads:
    key = h.split(": ", 1)[-1]
    if key not in seen:
        seen.add(key)
        uniq.append(h)
(mono / "_exit_headings.txt").write_text("\n".join(uniq), encoding="utf-8")
print("headings", len(uniq))
print("\n".join(uniq[:200]))

# Extract chapter 1-2 headings/text from original Word
docx_files = list(mono.glob("*.docx"))
print("\nDOCX FILES:")
for f in docx_files:
    if f.name.startswith("~$"):
        continue
    print(" -", f.name.encode("unicode_escape").decode(), f.stat().st_size)

# Prefer main monograph file
candidates = [f for f in docx_files if "اصلی" in f.name or "اصلي" in f.name]
if not candidates:
    candidates = [f for f in docx_files if f.name == "Last.docx"]
if not candidates:
    candidates = [f for f in docx_files if not f.name.startswith("~$")]

main = candidates[0]
print("MAIN WORD", main.name.encode("unicode_escape").decode())
w = Document(str(main))
lines = []
for para in w.paragraphs:
    t = para.text.strip()
    if not t:
        continue
    style = para.style.name if para.style else ""
    lines.append(f"[{style}] {t}")
(mono / "_main_word_extract.txt").write_text("\n".join(lines), encoding="utf-8")
print("paragraphs", len(lines))
# also extract Last.docx headings
for name in ["Last.docx", "Last_Mono_Expanded.docx"]:
    p = mono / name
    if not p.exists():
        continue
    ww = Document(str(p))
    hl = []
    for para in ww.paragraphs:
        t = para.text.strip()
        style = para.style.name if para.style else ""
        if "Heading" in style or t.startswith("Chapter") or t.startswith("فصل"):
            hl.append(f"[{style}] {t}")
    (mono / f"_headings_{name.replace('.docx','')}.txt").write_text("\n".join(hl), encoding="utf-8")
    print("saved headings for", name, len(hl))
