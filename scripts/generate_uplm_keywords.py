# -*- coding: utf-8 -*-
"""Standalone Keywords page for NAMS English monograph — system-specific terms only."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
MONO = next(p for p in ROOT.iterdir() if p.is_dir() and (p / "exit.pdf").exists())
OUT = MONO / "U_P_L_M_KEYWORDS.docx"

FONT = "Times New Roman"

# System-specific keywords (no global methodology brand names)
KEYWORDS = (
    "Nokta Academy Management System, student enrollment and fee management, "
    "teacher payroll and academy tax reporting, parent attendance visibility, "
    "multilingual academy portal"
)


def set_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    p = doc.add_paragraph()
    r = p.add_run("Nokta Academy Management System (NAMS)")
    set_run(r, size=14, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    p = doc.add_paragraph()
    r = p.add_run("Keywords — Standalone (System-Specific)")
    set_run(r, size=12, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # APA-style Keywords line
    p = doc.add_paragraph()
    r1 = p.add_run("Keywords: ")
    set_run(r1, size=12, italic=True)
    r2 = p.add_run(KEYWORDS)
    set_run(r2, size=12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.left_indent = Inches(0.5)
    pf.space_after = Pt(18)

    # Also list as five clear items for easy copy into abstract
    p = doc.add_paragraph()
    r = p.add_run("Keyword list (5 terms):")
    set_run(r, size=12, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    items = [
        "Nokta Academy Management System",
        "student enrollment and fee management",
        "teacher payroll and academy tax reporting",
        "parent attendance visibility",
        "multilingual academy portal",
    ]
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"{i}. {item}")
        set_run(r, size=12)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        p.paragraph_format.left_indent = Inches(0.5)

    doc.save(str(OUT))
    print("CREATED", str(OUT).encode("unicode_escape").decode())


if __name__ == "__main__":
    main()
