# -*- coding: utf-8 -*-
"""Standalone Declaration of Originality for English NAMS monograph (APA-aligned)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
MONO = next(p for p in ROOT.iterdir() if p.is_dir() and (p / "exit.pdf").exists())
OUT = MONO / "U_P_L_M_DECLARATION_OF_ORIGINALITY.docx"

FONT = "Times New Roman"
PROJECT_TITLE = "Nokta Academy Management System"


def set_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(doc, text, *, bold=False, italic=False, align="justify", size=12, space_after=12, first_indent=False, space_before=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, italic=italic)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.first_line_indent = Inches(0.5) if first_indent else Inches(0)
    return p


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    add_para(doc, "Kabul University", bold=True, align="center", size=14, space_after=0)
    add_para(doc, "Faculty of Computer Science", bold=True, align="center", size=12, space_after=0)
    add_para(doc, "Department of Software Engineering", bold=True, align="center", size=12, space_after=24)

    add_para(doc, "DECLARATION OF ORIGINALITY", bold=True, align="center", size=14, space_after=18)

    add_para(
        doc,
        f'I, ________________________________, Student ID ________________, declare that this '
        f'final-year project report titled “{PROJECT_TITLE}” is my original work, completed under the '
        f"academic requirements of the Bachelor of Science in Software Engineering at Kabul University.",
        first_indent=True,
        space_after=12,
    )

    add_para(
        doc,
        "I confirm that all sources, references, and quoted materials used in this report have been "
        "duly acknowledged and cited in accordance with the American Psychological Association (APA) "
        "7th edition referencing standard adopted for this monograph. Where the ideas, data, figures, "
        "or wording of other authors have been used, appropriate in-text citations and a complete "
        "reference list have been provided.",
        first_indent=True,
        space_after=12,
    )

    add_para(
        doc,
        "I further declare that this work has not been submitted previously, in whole or in part, for "
        "any other degree, diploma, or academic award at Kabul University or at any other institution. "
        "Any contribution of software tools, libraries, or institutional data used in the development of "
        "the Nokta Academy Management System has been used solely for academic purposes and is described "
        "honestly in the report.",
        first_indent=True,
        space_after=12,
    )

    add_para(
        doc,
        "I understand that failure to maintain academic integrity, including plagiarism or unacknowledged "
        "use of another person’s work, may result in academic penalties in accordance with the regulations "
        "of Kabul University.",
        first_indent=True,
        space_after=24,
    )

    add_para(doc, "Student’s Full Name: ________________________________", align="left", space_after=8)
    add_para(doc, "Student ID: ________________", align="left", space_after=8)
    add_para(doc, "Signature: ________________________", align="left", space_after=8)
    add_para(doc, "Date: ____ / ____ / ________", align="left", space_after=24)

    add_para(
        doc,
        "If the project was completed jointly, each author must complete a separate signed copy of this declaration.",
        italic=True,
        align="left",
        size=11,
        space_after=0,
    )

    doc.save(str(OUT))
    print("CREATED", str(OUT).encode("unicode_escape").decode())


if __name__ == "__main__":
    main()
