# -*- coding: utf-8 -*-
"""Standalone APA-style Abstract for English NAMS monograph (U_P_L_M)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
MONO = next(p for p in ROOT.iterdir() if p.is_dir() and (p / "exit.pdf").exists())
OUT = MONO / "U_P_L_M_ABSTRACT.docx"

FONT = "Times New Roman"

ABSTRACT = (
    "Private academies that still rely on paper files, spreadsheets, and informal messaging often face "
    "fragmented student records, delayed reporting, weak parent communication, and limited financial "
    "transparency. This monograph presents the design and implementation of the Nokta Academy Management "
    "System (NAMS), an integrated web and mobile-oriented platform developed to centralize academic, "
    "administrative, and financial operations for Nokta Academy and similar private learning centers. "
    "The study used an applied mixed-methods approach: a structured questionnaire completed by 20 academies "
    "informed feature priorities, while workflow and role analysis guided functional and non-functional "
    "requirements. Development followed Agile Scrum within a documented software life-cycle, delivered in "
    "three reviewed prototypes. Prototype 1 provided the core baseline for authentication, student and "
    "teacher management, classes and subjects, manual attendance, examinations and results, fee collection, "
    "and baseline payroll, without AI exam generation, online meeting attendance automation, or full "
    "government tax packaging. Prototype 2 added AI-assisted question generation, online class meeting "
    "links, and automatic parental attendance visibility after 18 of the 20 academies requested these "
    "features; two newly opened academies offered no specific opinion. Prototype 3 introduced quarterly "
    "and annual government tax reporting at Nokta Academy’s request. System design used layered architecture "
    "with Node.js/Express, MongoDB, and React, supported by use-case, sequence, and entity-relationship "
    "models and verified through unit, integration, and user-acceptance testing. Results indicate strong "
    "adoption intent and clear demand for integrated fee, attendance, examination, parent-communication, "
    "and tax-reporting capabilities. The project concludes that an iteratively delivered, academy-specific "
    "management system can reduce manual workload, improve information quality, and support compliance-"
    "oriented financial reporting within the stated scope and sample limitations."
)

KEYWORDS = (
    "academy management system, Agile Scrum, educational information systems, "
    "learning analytics readiness, private academy digitization"
)


def set_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    # Title of work (context only)
    p = doc.add_paragraph()
    r = p.add_run("Nokta Academy Management System (NAMS)")
    set_run(r, size=14, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    p = doc.add_paragraph()
    r = p.add_run("Final Year Project Monograph — Abstract (Standalone)")
    set_run(r, size=12, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # ABSTRACT heading (APA: centered, bold)
    p = doc.add_paragraph()
    r = p.add_run("Abstract")
    set_run(r, size=12, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

    # Abstract body — one paragraph, no first-line indent (APA)
    p = doc.add_paragraph()
    r = p.add_run(ABSTRACT)
    set_run(r, size=12)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Inches(0)
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(12)

    # Keywords — APA: indented, italic label
    p = doc.add_paragraph()
    r1 = p.add_run("Keywords: ")
    set_run(r1, size=12, italic=True)
    r2 = p.add_run(KEYWORDS)
    set_run(r2, size=12)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(0)
    pf.space_before = Pt(6)

    # Word count note (for author use; not part of formal abstract heading hierarchy)
    wc = len(ABSTRACT.split())
    p = doc.add_paragraph()
    r = p.add_run(f"[Author note — not for binding: abstract word count = {wc} words (target 200–300).]")
    set_run(r, size=10, italic=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    doc.save(str(OUT))
    print("CREATED", str(OUT).encode("unicode_escape").decode())
    print("WORDS", wc)


if __name__ == "__main__":
    main()
