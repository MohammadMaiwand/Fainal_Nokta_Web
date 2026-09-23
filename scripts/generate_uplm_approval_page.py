# -*- coding: utf-8 -*-
"""Standalone Approval Page for NAMS English monograph."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm

ROOT = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
MONO = next(p for p in ROOT.iterdir() if p.is_dir() and (p / "exit.pdf").exists())
OUT = MONO / "U_P_L_M_APPROVAL_PAGE.docx"

FONT = "Times New Roman"
PROJECT_TITLE = "Nokta Academy Management System"


def set_run(run, size=12, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_p(doc, text, *, bold=False, italic=False, align="center", size=12, before=0, after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, italic=italic)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Inches(0)
    return p


def set_cell_border(cell, **edges):
    """edges: top/bottom/left/right = size in eighths of a point, or None for nil."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        val = edges.get(edge)
        if val is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(val))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def fill_cell(cell, text, *, bold=False, size=11, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.15


def main():
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    add_p(doc, "Kabul University", bold=True, size=14, after=0)
    add_p(doc, "Faculty of Computer Science", bold=True, size=12, after=0)
    add_p(doc, "Department of Software Engineering", bold=True, size=12, after=24)

    add_p(doc, "APPROVAL PAGE", bold=True, size=14, before=6, after=18)

    add_p(
        doc,
        f'This project report titled “{PROJECT_TITLE}” is hereby approved as a satisfactory '
        f"report in partial fulfillment of the requirements for the Bachelor of Science in "
        f"Software Engineering.",
        align="justify",
        after=24,
    )

    # Approval table
    headers = ["Role", "Name", "Signature", "Date"]
    rows = [
        ["Project Supervisor", "________________________________", "_______________", "____________"],
        ["Department Head", "________________________________", "_______________", "____________"],
        ["External Examiner", "________________________________", "_______________", "____________"],
    ]

    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row — APA-like top/bottom line
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        fill_cell(cell, h, bold=True, size=11)
        set_cell_border(cell, top=12, bottom=12, left=None, right=None)

    for i, row_data in enumerate(rows):
        is_last = i == len(rows) - 1
        for j, val in enumerate(row_data):
            cell = table.rows[i + 1].cells[j]
            fill_cell(cell, val, bold=(j == 0), size=11)
            if is_last:
                set_cell_border(cell, top=None, bottom=12, left=None, right=None)
            else:
                set_cell_border(cell, top=None, bottom=None, left=None, right=None)

    widths = [Cm(4.2), Cm(6.5), Cm(3.5), Cm(3.0)]
    for row in table.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = w

    add_p(
        doc,
        "The undersigned certify that they have read this project report and that, in their opinion, "
        "it is satisfactory in scope and quality as a final-year project report for the degree of "
        "Bachelor of Science in Software Engineering.",
        align="justify",
        before=24,
        after=18,
    )

    add_p(doc, "Place of approval: Department of Software Engineering, Kabul University", align="left", after=6)
    add_p(doc, "Academic year: ____________________", align="left", after=18)

    add_p(
        doc,
        "Note. Please write the full academic title and name of each signatory clearly "
        "(for example, Dr. / Prof. followed by the full name).",
        italic=True,
        align="left",
        size=11,
        after=0,
    )

    doc.save(str(OUT))
    print("CREATED", str(OUT).encode("unicode_escape").decode())


if __name__ == "__main__":
    main()
