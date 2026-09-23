# -*- coding: utf-8 -*-
"""
Build U_P_L_M.docx:
- Chapters 1–2 copied from مونوگراف اصلی .docx (same design)
- Chapters 3–end follow Exit.pdf structure, content for NAMS
- Times New Roman; headings 14pt bold; body 12pt; APA-style tables/captions
"""
from __future__ import annotations

import shutil
from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(r"F:\Bcak_End_update_file-main\Nokta_V7 - Copy")
MONO = next(p for p in ROOT.iterdir() if p.is_dir() and (p / "exit.pdf").exists())
ASSETS = MONO / "_uplm_assets"
OUT = MONO / "U_P_L_M.docx"

FONT = "Times New Roman"
BODY_PT = 12
HEAD_PT = 14
LINE = 480  # double spacing (APA)


def set_run(run, size=BODY_PT, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_para(p, size=BODY_PT, bold=False, align="justify", space_after=8, first_indent=True):
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if first_indent and align == "justify":
        pf.first_line_indent = Inches(0.5)
    else:
        pf.first_line_indent = Inches(0)
    for r in p.runs:
        set_run(r, size=size, bold=bold)


def add_p(doc, text, *, size=BODY_PT, bold=False, align="justify", first_indent=True, space_after=8):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)
    style_para(p, size=size, bold=bold, align=align, space_after=space_after, first_indent=first_indent)
    return p


def add_h(doc, text, level=1):
    # User: titles 14 bold (use same size for H1/H2/H3; distinguish by numbering)
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size=HEAD_PT, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(18 if level == 1 else 12)
    pf.space_after = Pt(10)
    pf.first_line_indent = Inches(0)
    # outline level for TOC
    p.style = doc.styles["Heading %d" % min(level, 3)]
    # re-apply font after style
    for r in p.runs:
        set_run(r, size=HEAD_PT, bold=True)
    return p


def page_break(doc):
    doc.add_page_break()


def add_caption_table(doc, number, title):
    add_p(doc, f"Table {number}", size=BODY_PT, bold=True, align="left", first_indent=False, space_after=0)
    add_p(doc, title, size=BODY_PT, bold=False, align="left", first_indent=False, space_after=6)


def add_caption_figure(doc, number, title):
    add_p(doc, f"Figure {number}", size=BODY_PT, bold=True, align="left", first_indent=False, space_after=0)
    add_p(doc, title, size=BODY_PT, bold=False, align="left", first_indent=False, space_after=6)
    p = doc.add_paragraph()
    run = p.add_run(
        "Note. Generated from the Nokta Academy Management System (NAMS) project artifacts and evaluation data."
    )
    set_run(run, size=10, italic=True)
    style_para(p, size=10, align="left", first_indent=False, space_after=10)


def set_cell_text(cell, text, bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    set_run(run, size=size, bold=bold)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Inches(0)


def _set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        if val is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(val))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def apa_table(doc, headers, rows):
    """APA-ish: horizontal borders on header row (top+bottom) and last row (bottom only)."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for j, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], h, bold=True, size=11)
        _set_cell_borders(table.rows[0].cells[j], top=12, bottom=12, left=None, right=None)

    for i, row in enumerate(rows):
        is_last = i == len(rows) - 1
        for j, val in enumerate(row):
            set_cell_text(table.rows[i + 1].cells[j], val, bold=False, size=11)
            if is_last:
                _set_cell_borders(table.rows[i + 1].cells[j], top=None, bottom=12, left=None, right=None)
            else:
                _set_cell_borders(table.rows[i + 1].cells[j], top=None, bottom=None, left=None, right=None)

    # spacer after table
    add_p(doc, "", first_indent=False, space_after=4)
    return table


def add_image(doc, path: Path, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    return p


def figure_block(doc, path, number, title, before_paras, after_paras):
    for t in before_paras:
        add_p(doc, t)
    add_image(doc, path)
    add_caption_figure(doc, number, title)
    for t in after_paras:
        add_p(doc, t)


def table_block(doc, number, title, headers, rows, before_paras, after_paras):
    for t in before_paras:
        add_p(doc, t)
    add_caption_table(doc, number, title)
    apa_table(doc, headers, rows)
    for t in after_paras:
        add_p(doc, t)


def find_source_docx() -> Path:
    for f in MONO.glob("*.docx"):
        if f.name.startswith("~$"):
            continue
        if "اصلی" in f.name or "اصلي" in f.name:
            return f
    raise FileNotFoundError("مونوگراف اصلی .docx not found")


def configure_styles(doc: Document):
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3"):
        try:
            st = doc.styles[style_name]
        except KeyError:
            continue
        st.font.name = FONT
        st.font.size = Pt(HEAD_PT if "Heading" in style_name else BODY_PT)
        st.font.bold = "Heading" in style_name
        if st._element.rPr is not None and st._element.rPr.rFonts is not None:
            st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def build_chapter3(doc: Document):
    page_break(doc)
    add_h(doc, "CHAPTER 3: METHODOLOGY", 1)
    add_p(
        doc,
        "This chapter explains the research and development methods used to design and implement the "
        "Nokta Academy Management System (NAMS). Following the structural organization used in Exit.pdf, "
        "it presents requirements elicitation, population and sampling, data collection and analysis, "
        "functional and non-functional requirements, Agile/Scrum development, three prototype stages, "
        "system design artifacts, implementation, and testing. The chapter connects stakeholder needs "
        "identified for private academies in Afghanistan with the engineering decisions that shaped NAMS "
        "(Creswell & Creswell, 2018; Pressman & Maxim, 2020).",
    )

    # 3.1
    add_h(doc, "3.1 Research Methodology", 2)
    add_p(
        doc,
        "This study used an applied mixed-methods research design. Quantitative questionnaire data "
        "measured feature demand and adoption likelihood across academies, while qualitative workflow "
        "analysis and role analysis explained how registration, attendance, examinations, fees, payroll, "
        "and tax reporting are performed in daily operations (Creswell & Creswell, 2018). Mixed methods "
        "were appropriate because NAMS is not only a technical product; it is an institutional information "
        "system whose success depends on both measurable demand and contextual work practices "
        "(DeLone & McLean, 2003).",
    )
    add_p(
        doc,
        "Before the questionnaire was finalized, a preliminary feature set was derived from the literature "
        "reviewed in Chapter 2 and from inspection of related platforms (Moodle, Google Classroom, Canvas, "
        "and school ERP tools). Literature on learning analytics and educational data mining emphasized "
        "structured attendance, assessment, and communication data (Aldowah et al., 2019; Romero & Ventura, 2020). "
        "These sources produced an initial backlog that was then validated with academy stakeholders.",
    )

    add_h(doc, "3.1.1 Requirements Elicitation", 3)
    add_p(
        doc,
        "Requirements elicitation translated the problem statement in Chapter 1 into concrete system "
        "capabilities. The elicitation process combined (a) document and proposal review, (b) observation "
        "of academy workflows, (c) role-permission analysis for administrators, teachers, students, parents, "
        "owners, and branch managers, and (d) a structured questionnaire for multi-academy validation "
        "(International Organization for Standardization et al., 2018).",
    )
    add_p(
        doc,
        "The preliminary feature list included student enrollment, class and subject management, timetable "
        "management, manual attendance, examinations and results, fee collection, payroll estimation, "
        "messaging, dashboards, multilingual interfaces, and role-based access. Later backlog items—AI "
        "question generation, online class meeting links, automated parental attendance visibility, and "
        "government tax reporting—were treated as candidate enhancements pending survey and sponsor confirmation.",
    )
    add_p(
        doc,
        "Elicitation therefore produced two layers of requirements: a core operational layer required for "
        "daily academy management, and an enhancement layer validated by academy survey results and by "
        "Nokta Academy’s own compliance request. This separation later became the three-prototype plan "
        "described in Section 3.2.4.",
    )

    add_h(doc, "3.1.2 Population and Sampling", 3)
    add_p(
        doc,
        "The target population consisted of private academies and short-course centers that manage "
        "students, teachers, fees, and academic records in Afghanistan’s urban education market. Because "
        "a complete sampling frame of all academies was unavailable, purposive sampling was used to recruit "
        "academies already operating or recently established and willing to evaluate an integrated management "
        "platform (Creswell & Creswell, 2018).",
    )
    table_block(
        doc,
        "3.1",
        "Population, Sample, and Response Profile for NAMS Requirements Validation",
        ["Item", "Description", "Count / Value"],
        [
            ["Target population", "Private academies / course centers with academic–financial operations", "Urban academy segment"],
            ["Invited academies", "Academies contacted for questionnaire participation", "20"],
            ["Completed questionnaires", "Fully completed institutional questionnaires", "20"],
            ["Academies requesting Prototype-2 features", "AI exam support, online meeting link, auto parent attendance", "18"],
            ["Newly started academies (no specific opinion)", "Recently opened; limited operational history", "2"],
            ["Primary respondents", "Admin, owner, teacher, branch manager, accountant", "20 individuals"],
        ],
        [
            "Table 3.1 summarizes the sampling frame used for requirements validation. Twenty academies "
            "completed the questionnaire. Eighteen academies explicitly requested the advanced academic–parent "
            "features later packaged as Prototype 2, while two newly established academies did not offer a "
            "specific preference because their processes were still forming.",
            "The sample is small but information-rich for a final-year software engineering project. It is "
            "sufficient to prioritize features and to justify iterative delivery, but it is not intended as a "
            "nationally representative statistical inference (Brooke, 1996; Creswell & Creswell, 2018).",
        ],
        [
            "As shown in Table 3.1, the survey design treats each academy as an institutional case. This matches "
            "NAMS’s deployment model, in which configuration and reporting are academy-scoped rather than "
            "student-only. The 18-versus-2 split became a direct decision rule for moving AI examination support, "
            "online class links, and automated parental attendance into Prototype 2.",
        ],
    )

    add_h(doc, "3.1.3 Data Collection", 3)
    add_p(
        doc,
        "Data collection used a structured questionnaire (Appendix A) administered to academy managers and "
        "key staff. Items covered current tools, pain points, feature usefulness (Likert-type ratings), "
        "priority ranking, and likelihood of adoption (Likert, 1932). Complementary qualitative notes were "
        "collected during walkthroughs of fee, attendance, and payroll procedures at Nokta Academy.",
    )
    add_p(
        doc,
        "Questionnaire sections were mapped to design decisions: enrollment and fees mapped to finance "
        "modules; attendance and parent communication mapped to attendance policies and notifications; "
        "examination items mapped to exam and AI services; compliance items mapped to tax reporting. This "
        "traceability follows requirements engineering guidance that requirements remain testable and "
        "linked to source evidence (International Organization for Standardization et al., 2018).",
    )
    add_p(
        doc,
        "All responses were anonymized at academy level for reporting. No student personally identifiable "
        "grades were collected in the survey itself; evaluation used system test data and controlled "
        "user-acceptance scenarios instead.",
    )

    add_h(doc, "3.1.4 Data Analysis and Presentation", 3)
    add_p(
        doc,
        "Questionnaire responses were analyzed with descriptive statistics: frequencies, percentages, and "
        "mean priority scores. Results are presented in Chapter 4 using APA-style tables and figures. "
        "Qualitative workflow notes were coded into requirement themes (registration, attendance, exams, "
        "finance, tax, communication) and cross-checked against implemented modules (Miles et al., 2014).",
    )
    add_p(
        doc,
        "Where quantitative and qualitative evidence converged—for example, strong demand for fee transparency "
        "and parental attendance visibility—the feature was raised in the Product Backlog. Where evidence "
        "diverged—for example, newly opened academies without stable processes—the feature was deferred or "
        "kept optional through configuration flags.",
    )

    add_h(doc, "3.1.5 Functional Requirements", 3)
    table_block(
        doc,
        "3.2",
        "Selected Functional Requirements for NAMS",
        ["ID", "Requirement", "Priority", "Source", "Verification"],
        [
            ["FR-01", "Register and manage students with class, subject, teacher, and fee linkage", "High", "Workflow + survey", "CRUD/API tests"],
            ["FR-02", "Manage teachers, salary type (fixed/percentage/hybrid), and assignments", "High", "Payroll workflow", "Payroll tests"],
            ["FR-03", "Record attendance by session and report daily/period summaries", "High", "Survey", "Functional tests"],
            ["FR-04", "Create exams, record results, and publish to students/parents", "High", "Proposal", "Exam/result tests"],
            ["FR-05", "Collect fees, track balances, and link payments to payroll bases", "High", "Finance workflow", "Payment tests"],
            ["FR-06", "Enforce JWT authentication and role-based permissions", "High", "Security need", "Auth/RBAC tests"],
            ["FR-07", "Provide multilingual UI (English, Dari, Pashto) with RTL support", "Medium", "Context", "Localization tests"],
            ["FR-08", "Generate AI exam questions (Prototype 2)", "Medium", "18/20 academies", "AI endpoint tests"],
            ["FR-09", "Support online class meeting links and online attendance automation (Prototype 2)", "Medium", "18/20 academies", "Job/policy tests"],
            ["FR-10", "Produce quarterly/annual government tax reports (Prototype 3)", "High", "Nokta Academy request", "Tax service tests"],
        ],
        [
            "Functional requirements were written as verifiable statements so that each item could be traced "
            "to a module, API route, or automated test. Table 3.2 presents the core set used to drive the "
            "Product Backlog and Definition of Done for each sprint (Schwaber & Sutherland, 2020).",
            "High-priority items formed Prototype 1. FR-08 and FR-09 were explicitly deferred until survey "
            "confirmation. FR-10 was scheduled after Nokta Academy requested government tax compliance features.",
        ],
        [
            "Table 3.2 also shows verification methods. This linkage is essential in software engineering "
            "projects because unimplemented narrative requirements cannot demonstrate objective achievement "
            "(Pressman & Maxim, 2020; Sommerville, 2016).",
        ],
    )

    add_h(doc, "3.1.6 Non-Functional Requirements", 3)
    table_block(
        doc,
        "3.3",
        "Selected Non-Functional Requirements for NAMS",
        ["ID", "Category", "Requirement", "Verification"],
        [
            ["NFR-01", "Performance", "Common CRUD and dashboard queries respond under normal LAN load", "Manual/API timing"],
            ["NFR-02", "Security", "JWT sessions, password hashing, RBAC, CSRF protections on sensitive routes", "Security tests"],
            ["NFR-03", "Usability", "Role-focused menus; readable light/dark UI; multilingual labels", "UAT/SUS-style feedback"],
            ["NFR-04", "Maintainability", "Modular Express routes/services and feature-based React structure", "Code review"],
            ["NFR-05", "Reliability", "Financial and tax calculations are deterministic and auditable", "Unit tests"],
            ["NFR-06", "Compatibility", "Modern desktop browsers; Flutter/PWA mobile access paths", "Device checks"],
            ["NFR-07", "Localizability", "EN/FA/PS resources and Jalali/Hijri-aware finance filters", "Locale tests"],
        ],
        [
            "Non-functional requirements define quality attributes that determine whether NAMS is usable in "
            "real academies. Security and reliability are especially important because the system stores "
            "student records, fees, salaries, and tax figures (International Organization for Standardization "
            "& International Electrotechnical Commission, 2011).",
            "Usability and localizability address adoption barriers identified in Chapter 2, including "
            "multilingual communities and mixed digital literacy among staff and parents (Davis, 1989; "
            "Venkatesh et al., 2003).",
        ],
        [
            "Table 3.3 was used during sprint reviews as a quality checklist. A feature was not marked Done "
            "when it worked only for one role or one language; it had to satisfy the relevant NFR evidence.",
        ],
    )

    add_h(doc, "3.1.7 Data Requirements", 3)
    add_p(
        doc,
        "NAMS requires structured records for users, roles, branches, students, teachers, parents/families, "
        "classes, subjects, courses, enrollments, timetables, attendance and attendance policies, exams, "
        "questions, results, payments, expenses, salary records, tax settings, messages, notifications, "
        "and audit logs. Sensitivity is high because records combine personal, academic, and financial data.",
    )
    add_p(
        doc,
        "Data requirements also anticipate analytics. Consistent attendance, assessment, and fee histories "
        "are prerequisites for future learning-analytics and AI support services (Aldowah et al., 2019; "
        "Siemens, 2013). Prototype 1 therefore prioritized clean relational references even when advanced "
        "AI features were not yet enabled.",
    )
    add_p(
        doc,
        "MongoDB was selected as the persistence engine, with ObjectId references used to preserve "
        "relationships needed for payroll, tax, and reporting joins (Chodorow, 2013; Connolly & Begg, 2015).",
    )

    # 3.2 Development methodology
    add_h(doc, "3.2 Development Methodology", 2)
    add_p(
        doc,
        "NAMS was developed using an Agile Scrum approach embedded in a lightweight software development "
        "life cycle (SDLC) envelope. Agile practices managed changing academy priorities; SDLC stages "
        "(analysis, design, implementation, testing, deployment readiness) preserved academic documentation "
        "discipline (Pressman & Maxim, 2020; Schwaber & Sutherland, 2020).",
    )

    add_h(doc, "3.2.1 Agile Development Approach", 3)
    add_p(
        doc,
        "Agile was selected because academy stakeholders could not freeze every requirement at the start. "
        "Fee rules, parental expectations, online-class habits, and tax reporting obligations clarified "
        "only after early demos. Agile’s iterative delivery reduced the risk of building a polished but "
        "misaligned system (Dikert et al., 2016).",
    )
    add_p(
        doc,
        "The Agile practices actually used were: Product Backlog prioritization, time-boxed sprints, "
        "increment demos (Sprint Review), a working Definition of Done, and backlog refinement after "
        "questionnaire feedback. Formal daily Scrum ceremonies were simplified for a small student team, "
        "but the inspect-and-adapt loop remained central (Schwaber & Sutherland, 2020).",
    )
    figure_block(
        doc,
        ASSETS / "fig_agile_scrum.png",
        "3.0",
        "Agile Scrum cycle applied to NAMS prototype stages",
        [
            "Figure 3.0 summarizes how Agile cycles map onto NAMS delivery. The Product Backlog held core "
            "operations first, then survey-driven enhancements, then sponsor-requested tax compliance. Each "
            "cycle ended with review evidence rather than only code completion.",
            "This figure is placed before the prototype subsections so that readers can see why Prototype 1 "
            "intentionally omitted some popular features. Omission was a prioritization decision, not an "
            "oversight, and it matches Agile value delivery under constrained academic timelines.",
        ],
        [
            "After Figure 3.0, the narrative proceeds from process to product: Scrum organization, iteration "
            "planning, and the three reviewed prototypes that operationalize the backlog.",
        ],
    )

    add_h(doc, "3.2.2 Scrum-Based Project Organization", 3)
    add_p(
        doc,
        "Roles were adapted to the academic project setting. The Product Owner role was shared between the "
        "student team and Nokta Academy management for priority decisions. Developers implemented backend, "
        "web, and mobile slices. Testing combined automated unit checks with stakeholder UAT during sprint "
        "reviews (Schwaber & Sutherland, 2020).",
    )
    add_p(
        doc,
        "Work items were sliced vertically where possible (for example, student registration through API, "
        "UI, and permission checks in one increment) to avoid long integration delays. Financial algorithms "
        "received extra unit-test attention because calculation defects harm trust quickly "
        "(DeLone & McLean, 2003).",
    )
    add_p(
        doc,
        "Documentation artifacts for the monograph—use cases, sequence diagrams, architecture, ERD, and "
        "test tables—were updated at the end of major increments so that the report remained synchronized "
        "with the running system.",
    )

    add_h(doc, "3.2.3 Development Iterations", 3)
    add_p(
        doc,
        "Iterations followed a staged capability model: (1) authentication, users, and academic masters; "
        "(2) attendance, exams, and results; (3) payments and payroll baseline; (4) communication and "
        "dashboards; (5) Prototype-2 enhancements; (6) Prototype-3 tax compliance. Each iteration produced "
        "a demonstrable increment for review.",
    )
    add_p(
        doc,
        "This ordering matches dependency logic: payroll needs payments and teacher assignments; tax "
        "reports need stable finance records; parental attendance automation needs attendance policies and "
        "notification channels. Agile iteration did not remove dependencies; it exposed them earlier "
        "(Sommerville, 2016).",
    )
    add_p(
        doc,
        "Regression checks after later iterations protected earlier increments. For example, adding AI "
        "question generation could not break manual exam scoring paths used by teachers who prefer paper "
        "assessments.",
    )

    add_h(doc, "3.2.4 Prototype Development and User Review", 3)
    add_p(
        doc,
        "Prototype development was organized in three reviewed stages. Prototype 1 delivered the integrated "
        "operational baseline. Prototype 2 added features demanded by 18 of 20 surveyed academies. "
        "Prototype 3 added quarterly and annual government tax reporting requested by Nokta Academy. "
        "Each stage was reviewed with stakeholders before the next backlog package was locked.",
    )

    # Prototype 1
    add_h(doc, "3.2.4.1 Prototype 1 — Core Integrated Baseline", 3)
    add_p(
        doc,
        "Prototype 1 implemented the minimum integrated platform required to replace fragmented "
        "spreadsheet and paper workflows: authentication and RBAC, student and teacher management, "
        "classes/subjects/timetable, manual attendance, manual examinations and results, fee collection, "
        "baseline payroll estimation, messaging/announcements, and role-based dashboards.",
    )
    add_p(
        doc,
        "Explicitly excluded from Prototype 1 were: (a) AI-assisted examination question generation, "
        "(b) Google Meet / online class meeting-link workflows as a first-class attendance channel, and "
        "(c) automatic attendance marking for parental monitoring. These exclusions kept the first "
        "increment focused on data integrity and core CRUD reliability before automation complexity "
        "was introduced (Pressman & Maxim, 2020).",
    )
    figure_block(
        doc,
        ASSETS / "fig_3_1_proto1_dashboard.png",
        "3.1",
        "Prototype 1 admin dashboard for core academy operations",
        [
            "Figure 3.1 shows the Prototype 1 administrative dashboard used during the first review. The "
            "screen emphasizes enrollment, manual attendance, fee status, and operational reports rather "
            "than AI or online-class automation.",
            "Reviewers used this dashboard to verify that key KPIs were reachable in one navigation hop, "
            "which supports perceived ease of use in technology acceptance models (Davis, 1989).",
        ],
        [
            "The dashboard in Figure 3.1 confirmed that Prototype 1 already reduced tool-switching among "
            "admin staff. Remaining complaints concerned examination authoring speed and parent visibility—"
            "exactly the gaps later addressed in Prototype 2.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_2_proto1_students.png",
        "3.2",
        "Prototype 1 student, class, teacher, and timetable management screens",
        [
            "Figure 3.2 presents the academic master-data screens of Prototype 1. Student records are linked "
            "to class, subject, teacher, and fee fields so that later payment and payroll calculations have "
            "stable references.",
            "Without these linkages, financial modules would depend on manual re-entry and would recreate "
            "the fragmentation problem identified in Chapter 1 (Connolly & Begg, 2015).",
        ],
        [
            "Stakeholder review of Figure 3.2 focused on validation rules (required fields, duplicate IDs, "
            "and assignment consistency). Fixes from that review entered the backlog before Prototype 2 work began.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_3_proto1_attendance_exam.png",
        "3.3",
        "Prototype 1 manual attendance and basic examination workflow",
        [
            "Figure 3.3 documents the Prototype 1 attendance and examination path. Teachers mark attendance "
            "manually, and exams are created and scored without AI generation. Parents may view published "
            "results, but they do not receive automated online-attendance marks in this stage.",
            "This design was intentional: automation was postponed until policy flags, notification channels, "
            "and survey demand were confirmed.",
        ],
        [
            "Figure 3.3 therefore represents a controlled baseline against which Prototype 2 enhancements "
            "can be compared in Chapter 4 results.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_4_proto1_finance.png",
        "3.4",
        "Prototype 1 fee collection and baseline payroll view",
        [
            "Figure 3.4 shows Prototype 1 finance capabilities: tuition collection, balances, expense "
            "recording, and baseline payroll estimation. Government quarterly/annual tax packages are not "
            "yet included; only basic tax estimation appears where needed for salary net pay.",
            "Finance reviewers accepted this baseline for day-to-day fee operations but requested a later "
            "compliance package, which became Prototype 3.",
        ],
        [
            "Together, Figures 3.1–3.4 define the Prototype 1 contract: an integrated academy MIS without "
            "AI exam authoring, without online-meeting attendance automation, and without full government "
            "tax reporting.",
        ],
    )

    # Prototype 2
    add_h(doc, "3.2.4.2 Prototype 2 — Survey-Driven Academic and Parent Enhancements", 3)
    add_p(
        doc,
        "Prototype 2 was authorized only after questionnaire analysis. Of 20 academies that completed the "
        "instrument, 18 stated that they needed AI-supported examination question preparation, online class "
        "meeting support, and automatic attendance visibility for parents. Two newly opened academies did "
        "not provide a specific opinion because their teaching and parent-communication routines were not "
        "yet stable.",
    )
    add_p(
        doc,
        "Based on that 18/20 demand signal, Prototype 2 added: (1) AI question generation for MCQ, "
        "true/false, and descriptive items with teacher review; (2) timetable online meeting links "
        "(Google Meet or equivalent URL) for remote/hybrid sessions; and (3) attendance-policy-driven "
        "automatic online attendance marking with parent notifications and parent portal reports.",
    )
    add_p(
        doc,
        "These features were implemented behind permissions and configuration flags so that academies "
        "that were not ready could keep Prototype 1 behavior. Optional activation is consistent with "
        "Agile incremental adoption and reduces forced process change (Dikert et al., 2016).",
    )
    figure_block(
        doc,
        ASSETS / "fig_3_5_proto2_ai_exam.png",
        "3.5",
        "Prototype 2 AI-assisted examination question generation",
        [
            "Figure 3.5 illustrates the AI examination workspace introduced in Prototype 2. Teachers specify "
            "topic, difficulty, language, and question type; generated drafts are editable before publishing.",
            "The design keeps humans in control: AI accelerates authoring, but academic responsibility "
            "remains with the teacher (Romero & Ventura, 2020).",
        ],
        [
            "During review of Figure 3.5, teachers reported time savings in first-draft preparation, which "
            "aligns with the survey priority for faster exam setup.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_6_proto2_meet_attendance.png",
        "3.6",
        "Prototype 2 online class link and automatic parental attendance visibility",
        [
            "Figure 3.6 shows how Prototype 2 connects timetable online links with attendance automation. "
            "When a session is configured for online delivery and auto-marking is enabled, the system can "
            "create online attendance records and notify linked parents.",
            "Parents do not mark attendance themselves; they receive visibility and alerts. This matches "
            "the academy request for accountability without transferring control away from teachers/admin.",
        ],
        [
            "Figure 3.6 closes the communication gap identified in Chapter 1: parents previously relied on "
            "informal messaging to learn whether a child attended an online session.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_7_proto2_survey_impact.png",
        "3.7",
        "Prototype 2 feature package justified by the 20-academy questionnaire",
        [
            "Figure 3.7 summarizes the survey justification for Prototype 2. Eighteen academies requested "
            "the package; two newly started academies offered no specific opinion. The figure is retained "
            "in the methodology chapter because it is a process artifact as much as a result artifact.",
            "Presenting the decision evidence here maintains requirements traceability from elicitation to "
            "implementation (International Organization for Standardization et al., 2018).",
        ],
        [
            "Chapter 4 later expands the same survey into fuller descriptive charts. Figure 3.7 is the "
            "decision snapshot used by the Product Owner to approve Sprint scope for Prototype 2.",
        ],
    )

    # Prototype 3
    add_h(doc, "3.2.4.3 Prototype 3 — Government Tax Reporting (Quarterly and Annual)", 3)
    add_p(
        doc,
        "Prototype 3 was driven by Nokta Academy’s own institutional request rather than by the multi-academy "
        "feature poll. After daily fee and payroll operations stabilized, academy management required "
        "structured support for government tax obligations on a quarterly and annual basis, including "
        "academy-level business receipt tax (BRT) treatment and employee salary tax summaries.",
    )
    add_p(
        doc,
        "The implemented package includes Hijri-aware filters for month, quarter (Q1–Q4), and year; "
        "employee wage-tax brackets; academy tuition-based BRT calculation; annual packaging of quarterly "
        "totals; PDF/print export; and audit logging of report generation. Nonprofit education treatment "
        "for annual income tax is reflected according to the academy’s documented local compliance rules "
        "encoded in the finance services.",
    )
    add_p(
        doc,
        "Prototype 3 therefore extends NAMS from an operational MIS into a compliance-supporting platform. "
        "This stage demonstrates why Agile delivery mattered: tax scope was not fully clear at project "
        "start, yet the modular finance architecture allowed addition without rewriting academic modules "
        "(Bass et al., 2012).",
    )
    figure_block(
        doc,
        ASSETS / "fig_3_8_proto3_tax_dashboard.png",
        "3.8",
        "Prototype 3 government tax dashboard requested by Nokta Academy",
        [
            "Figure 3.8 presents the Prototype 3 tax dashboard. Administrators and owners can open quarterly "
            "and annual views without exporting raw payment tables into external spreadsheets.",
            "This directly responds to the sponsor’s compliance request and reduces manual reconciliation "
            "effort that previously risked calculation drift.",
        ],
        [
            "Review comments on Figure 3.8 emphasized clear separation between employee salary tax and "
            "academy BRT so that users do not confuse the two obligations.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_9_proto3_quarterly.png",
        "3.9",
        "Prototype 3 quarterly tax report filters and totals",
        [
            "Figure 3.9 shows quarterly reporting controls. Users select Hijri year and quarter, then generate "
            "totals for tuition base, BRT, and related payroll withholdings.",
            "Quarterly cadence matches the academy’s declared filing rhythm and supports incremental "
            "verification before annual consolidation.",
        ],
        [
            "The quarterly view in Figure 3.9 also aids branch comparison when more than one branch "
            "contributes tuition in the same period.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_10_proto3_annual.png",
        "3.10",
        "Prototype 3 annual tax package and export workspace",
        [
            "Figure 3.10 illustrates the annual package that consolidates four quarters into a year-to-date "
            "compliance set. Exports are printable and archivable for ownership review.",
            "Annual packaging was a sponsor-critical outcome of Prototype 3 and is evaluated again in "
            "Chapters 4 and 5 as a practical contribution beyond generic LMS features.",
        ],
        [
            "Figure 3.10 completes the prototype story: Prototype 1 stabilized operations, Prototype 2 "
            "answered multi-academy academic/parent demand, and Prototype 3 answered Nokta Academy’s "
            "government tax mandate.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_3_11_proto3_payroll_tax.png",
        "3.11",
        "Prototype 3 payroll–tax integration overview",
        [
            "Figure 3.11 connects salary records to tax services and finance entries. Gross, tax, and net "
            "values remain explainable for each teacher or staff member included in a period.",
            "This transparency is essential for trust in information systems that affect livelihood and "
            "compliance simultaneously (DeLone & McLean, 2003).",
        ],
        [
            "With Figures 3.8–3.11 reviewed, the methodology transitions from prototype evidence to formal "
            "system design models in Section 3.3.",
        ],
    )

    # 3.3 System Design
    add_h(doc, "3.3 System Design", 2)
    add_p(
        doc,
        "System design translates validated requirements into structural models: use cases, sequence "
        "diagrams, architecture, and database design. UML-style artifacts communicate responsibilities "
        "among actors and components before further implementation detail is discussed "
        "(Fowler, 2004; Bass et al., 2012).",
    )

    add_h(doc, "3.3.1 Use Case Model", 3)
    figure_block(
        doc,
        ASSETS / "fig_3_12_use_case.png",
        "3.12",
        "Use case diagram for the Nokta Academy Management System",
        [
            "Figure 3.12 presents the NAMS use case model. Primary actors are Student, Parent, Teacher, "
            "Admin, Super Admin, and Owner. Shared use cases include managing students and teachers, "
            "recording attendance, managing exams (including AI-assisted authoring in Prototype 2), "
            "collecting fees, payroll and tax reporting, messaging, online class links, and dashboards.",
            "The diagram is intentionally role-centric because authorization boundaries are part of the "
            "product, not an afterthought. Each actor sees only the use cases allowed by RBAC policy.",
            "Compared with a generic LMS use case set, NAMS includes fee, payroll, tax, and parent "
            "attendance visibility as first-class academic-management concerns.",
        ],
        [
            "Figure 3.12 guides permission matrix design in the backend and protected routes in the web "
            "client. If a use case appears for Parent, corresponding API scopes must exist; otherwise the "
            "model would over-promise capability.",
            "The next subsection details dynamic behavior for three critical flows using sequence diagrams.",
        ],
    )

    add_h(doc, "3.3.2 Sequence Diagrams", 3)
    figure_block(
        doc,
        MONO / "Figure_3_13_Student_Enrollment_and_Fee_Payment_Sequence.png",
        "3.13",
        "Sequence diagram for student enrollment and fee payment",
        [
            "Figure 3.13 models enrollment and fee payment. An administrator creates or updates a student, "
            "links class/subject/teacher data, and records payments that update balances and feed payroll "
            "bases where percentage salary rules apply.",
            "The sequence clarifies that payment is not an isolated transaction; it emits financial side "
            "effects consumed by later payroll and tax services.",
            "This design choice is central to NAMS’s claim of financial transparency relative to "
            "spreadsheet-based academies.",
        ],
        [
            "Reading Figure 3.13 beside Prototype 1 screens shows that the enrollment–payment path was "
            "available before AI and tax packages, establishing the data spine of the system.",
        ],
    )
    figure_block(
        doc,
        MONO / "Figure_3_14_Attendance_and_Examination_Sequence.png",
        "3.14",
        "Sequence diagram for attendance and examination processing",
        [
            "Figure 3.14 shows attendance marking and examination processing. In Prototype 1, marking is "
            "manual. In Prototype 2, online sessions may trigger automated online attendance records and "
            "parent notifications when policy allows.",
            "Examination flow includes creating an exam, optionally generating AI questions, publishing, "
            "collecting results, and releasing grades to students and parents.",
            "The sequence therefore spans both baseline and enhanced behaviors without requiring separate "
            "unrelated systems.",
        ],
        [
            "Figure 3.14 also supports testing design: unit tests cover scoring rules, while integration "
            "tests cover publish-and-notify paths.",
        ],
    )
    figure_block(
        doc,
        MONO / "Figure_3_15_Payroll_Tax_and_Financial_Report_Sequence.png",
        "3.15",
        "Sequence diagram for payroll, tax, and financial reporting",
        [
            "Figure 3.15 describes payroll and tax reporting interactions introduced fully in Prototype 3. "
            "Salary records are computed from configured rules and payment bases; tax services apply "
            "employee brackets and academy BRT logic; reports are filtered by Hijri month, quarter, or year.",
            "Owners and admins export printable summaries for compliance review, while audit logs retain "
            "who generated which package.",
            "The sequence demonstrates that compliance reporting reuses operational data rather than "
            "asking staff to retype totals into external templates.",
        ],
        [
            "Figures 3.13–3.15 collectively document the three most risk-sensitive business paths in NAMS: "
            "money in, academic evidence, and money/tax out.",
        ],
    )

    add_h(doc, "3.3.3 Architectural Design", 3)
    figure_block(
        doc,
        MONO / "Figure_3_16_NAMS_Architecture_Diagram.png",
        "3.16",
        "Layered architecture of NAMS (web, API, data, and supporting services)",
        [
            "Figure 3.16 presents the NAMS architecture. Clients include the React web application, PWA "
            "access, and Flutter mobile paths. The Express/Node.js API layer hosts authentication, "
            "authorization, domain services, and scheduled jobs. MongoDB stores documents, while supporting "
            "services cover email/SMS notifications, AI providers, and report generation.",
            "Layering isolates UI churn from business rules and persistence, improving maintainability "
            "(Bass et al., 2012).",
            "Security cross-cuts the API through JWT validation, branch scoping, permission middleware, "
            "and audit logging.",
        ],
        [
            "The architecture in Figure 3.16 explains how Prototype 2 jobs (auto attendance) and Prototype 3 "
            "tax services could be added as modules without redesigning the client stack.",
        ],
    )
    table_block(
        doc,
        "3.4",
        "Technology Stack and Rationale for NAMS",
        ["Layer", "Technology", "Rationale"],
        [
            ["Web client", "React 18, Vite, TypeScript, Tailwind", "Responsive SPA with modular features"],
            ["Mobile", "Flutter / PWA", "Student/teacher/parent access on phones"],
            ["API", "Node.js, Express, TypeScript", "Non-blocking I/O and modular routes"],
            ["Database", "MongoDB, Mongoose", "Flexible documents with ObjectId references"],
            ["AuthZ/AuthN", "JWT, bcrypt, RBAC permissions", "Secure role-scoped access"],
            ["Jobs", "node-cron attendance/notification jobs", "Automation for Prototype 2 policies"],
            ["AI", "OpenAI-compatible / Groq / Gemini providers", "Optional exam authoring support"],
            ["Reporting", "PDF/print pipelines, dashboards", "Operational and tax outputs"],
        ],
        [
            "Table 3.4 records stack choices that support the architecture in Figure 3.16. Selections "
            "favor rapid iteration for a small team while remaining familiar in contemporary full-stack "
            "practice (Sommerville, 2016).",
            "AI providers are optional dependencies; Prototype 1 remains operable if AI keys are absent.",
        ],
        [
            "Stack rationale in Table 3.4 is revisited in Chapter 5 when comparing NAMS with LMS-centric "
            "platforms that lack native payroll/tax modules.",
        ],
    )

    add_h(doc, "3.3.4 Database Design", 3)
    figure_block(
        doc,
        MONO / "Figure_3_17_NAMS_Entity_Relationship_Diagram.png",
        "3.17",
        "Entity-relationship overview of core NAMS collections",
        [
            "Figure 3.17 provides an ER-style overview of core entities: users/roles, students, teachers, "
            "classes, subjects, enrollments, attendance, exams/results, payments, salary records, and "
            "messages. Although MongoDB is document-oriented, references preserve relational integrity for "
            "reporting (Chodorow, 2013).",
            "The ER view helps justify why payment rows store student and contextual links used later by "
            "payroll and tax aggregations.",
            "Soft-delete fields and timestamps support auditability without physically destroying history "
            "needed for quarterly reports.",
        ],
        [
            "Figure 3.17 should be read with Figure 3.18, which shows a complementary schema view of "
            "document shapes used in implementation.",
        ],
    )
    figure_block(
        doc,
        MONO / "Figure_3_18_NAMS_Core_Relational_Document_Schema.png",
        "3.18",
        "Core relational-document schema used by NAMS services",
        [
            "Figure 3.18 details key fields and relationships among operational collections. It is the "
            "schema companion to the ER overview and is useful for developers tracing service queries.",
            "AttendancePolicy fields such as onlineAutoMarkEnabled appear here as configuration that "
            "enables Prototype 2 behavior without schema forks.",
            "Tax and salary fields show how Prototype 3 reads operational collections rather than "
            "maintaining an unrelated parallel ledger.",
        ],
        [
            "With architecture and data design established, Section 3.4 summarizes how the design was "
            "implemented in engineering phases.",
        ],
    )

    # 3.4 Implementation
    add_h(doc, "3.4 Implementation", 2)
    add_p(
        doc,
        "Implementation followed the architecture and prototype plan. Backend modules were organized by "
        "domain routes and services; the web client used feature folders and shared API clients; mobile "
        "paths consumed the same REST API. Continuous integration lint/build checks reduced integration "
        "surprises as sprints advanced.",
    )

    add_h(doc, "3.4.1 Development Environment", 3)
    add_p(
        doc,
        "The development environment used Node.js 18+, MongoDB locally, React with Vite, TypeScript across "
        "stacks, Git for version control, and environment files for secrets (never committed). Optional "
        "services—SMTP, SMS, Stripe, AI providers—could run in console/mock mode for campus demos.",
    )
    add_p(
        doc,
        "Windows offline launcher scripts support LAN demonstration without public internet, which matters "
        "for academies with unstable connectivity (a constraint noted in Chapter 2).",
    )
    add_p(
        doc,
        "Code quality gates included linting, unit tests for finance/tax calculations, and build verification "
        "for frontend and backend packages.",
    )

    add_h(doc, "3.4.2 Implementation Phases", 3)
    add_p(
        doc,
        "Phase A delivered identity, RBAC, students, teachers, classes, and subjects. Phase B delivered "
        "attendance and examinations. Phase C delivered payments and baseline payroll. Phase D delivered "
        "messaging, notifications, and dashboards (Prototype 1 complete). Phase E delivered AI exams, "
        "online links, and parental auto-attendance (Prototype 2). Phase F delivered quarterly/annual tax "
        "packages (Prototype 3).",
    )
    add_p(
        doc,
        "Each phase ended with a review checklist mapped to FR/NFR IDs. Incomplete checklist items returned "
        "to the backlog rather than being silently skipped (Schwaber & Sutherland, 2020).",
    )
    add_p(
        doc,
        "This phased evidence is later summarized in Chapter 4 as engineering results distinct from "
        "questionnaire results.",
    )

    add_h(doc, "3.4.3 Key Modules Implemented", 3)
    add_p(
        doc,
        "Key modules include authentication/session management, branch-aware masters, enrollment-linked "
        "students, attendance and attendance policies, exams/questions/results, AI generation endpoints, "
        "payments and finance entries, payroll and employee/academy tax services, messages/notifications, "
        "reports/printing, audit logs, and role dashboards. Enterprise extensions exist for transcripts, "
        "assignments, certificates, and security events where enabled.",
    )
    add_p(
        doc,
        "Frontend resource configuration drives many CRUD screens, reducing duplication while still allowing "
        "specialized wizards for payments and exams. This pattern sped Prototype 1 delivery without freezing "
        "UX customization for high-traffic flows.",
    )
    add_p(
        doc,
        "Background jobs implement Prototype 2 automation (online attendance auto-mark and related "
        "notifications) and are controlled by configuration so production behavior remains explicit.",
    )

    add_h(doc, "3.4.4 Challenges and Solutions", 3)
    table_block(
        doc,
        "3.5",
        "Implementation Challenges and Solutions in NAMS",
        ["Challenge", "Impact", "Solution"],
        [
            ["Changing feature priorities after demos", "Scope risk", "Agile backlog + three prototype gates"],
            ["Percentage payroll tied to payments", "Calculation disputes", "Deterministic services + unit tests"],
            ["Multilingual RTL UI", "Layout defects", "Locale resources + RTL testing"],
            ["Online attendance trust", "Parent complaints", "Policy flags + audit of auto marks"],
            ["Tax rule clarity", "Compliance risk", "Sponsor workshops + Prototype 3 dedicated package"],
            ["AI question quality", "Teacher distrust", "Editable drafts + human publish step"],
        ],
        [
            "Table 3.5 records the main engineering challenges. Most were organizational/requirements "
            "issues rather than unsolvable technical blockers—another reason Agile suited this project "
            "(Dikert et al., 2016).",
            "Tax clarity required sponsor workshops because legal interpretation cannot be guessed from "
            "generic LMS literature.",
        ],
        [
            "The solutions column in Table 3.5 feeds directly into the testing strategy: every major "
            "mitigation has a corresponding verification activity in Section 3.5.",
        ],
    )

    # 3.5 Testing
    add_h(doc, "3.5 Testing", 2)
    add_p(
        doc,
        "Testing combined unit, integration, and user acceptance testing. Financial and tax functions "
        "received the strictest unit coverage because defects there undermine institutional trust. "
        "Integration tests exercised cross-module flows such as enroll–pay–payroll and exam–result–notify. "
        "UAT validated prototype goals with academy reviewers (Sommerville, 2016).",
    )

    add_h(doc, "3.5.1 Unit Testing", 3)
    table_block(
        doc,
        "3.6",
        "Sample Unit Test Cases for Core NAMS Calculations and Rules",
        ["Test ID", "Area", "Scenario", "Expected Result"],
        [
            ["UT-01", "Payroll", "Percentage salary on assigned student payments", "Gross = base × percentage"],
            ["UT-02", "Employee tax", "Wage falls in progressive bracket", "Tax matches bracket rule"],
            ["UT-03", "BRT", "Quarterly tuition aggregate", "BRT = configured % of taxable tuition"],
            ["UT-04", "Attendance policy", "onlineAutoMarkEnabled = false", "No auto online marks created"],
            ["UT-05", "Attendance policy", "onlineAutoMarkEnabled = true for online session", "Auto mark + parent notify path"],
            ["UT-06", "AI exam", "Generate MCQ then edit before publish", "Draft editable; publish uses final text"],
            ["UT-07", "RBAC", "Student token calls payroll report", "Denied"],
            ["UT-08", "Validation", "Duplicate enrollment key", "Rejected with clear error"],
        ],
        [
            "Table 3.6 samples unit tests that protect Prototype 1–3 business rules. UT-04 and UT-05 "
            "explicitly lock the Prototype 1 versus Prototype 2 attendance behaviors.",
            "UT-02 and UT-03 protect Prototype 3 compliance math from silent regression when other modules change.",
        ],
        [
            "Unit tests in Table 3.6 were executed during CI and before sprint reviews that claimed a "
            "prototype increment complete.",
        ],
    )

    add_h(doc, "3.5.2 Integration Testing", 3)
    add_p(
        doc,
        "Integration testing verified multi-module sequences: student registration to payment posting; "
        "manual attendance to report totals; exam publish to parent-visible results; online session to auto "
        "attendance notification; salary calculation to quarterly tax export. Failures often appeared at "
        "boundary assumptions (missing branchId, unpublished exam, disabled job flags).",
    )
    add_p(
        doc,
        "API contract checks confirmed that frontend feature modules and backend routes agreed on payload "
        "shapes for payments, exams, and tax filters. This reduced UI-only “success” that hid backend rejects.",
    )
    add_p(
        doc,
        "Integration evidence is summarized qualitatively in Chapter 4 engineering results rather than as "
        "raw CI logs, which are retained in the repository.",
    )

    add_h(doc, "3.5.3 User Acceptance Testing", 3)
    add_p(
        doc,
        "UAT was conducted at the end of each prototype gate. Prototype 1 UAT asked whether staff could "
        "run daily enrollment, attendance, exams, and fees without spreadsheets. Prototype 2 UAT asked "
        "whether AI drafting, online links, and parental auto-attendance were understandable and trustworthy. "
        "Prototype 3 UAT asked whether quarterly/annual tax outputs matched management expectations.",
    )
    add_p(
        doc,
        "Feedback instruments combined task success observation with short satisfaction prompts inspired by "
        "usability evaluation practice (Brooke, 1996; Nielsen, 1994). Formal large-sample SUS scoring was "
        "not claimed; instead, acceptance decisions were recorded per prototype gate.",
    )
    add_p(
        doc,
        "UAT outcomes indicated acceptance of the Prototype 1 baseline, strong interest in Prototype 2 among "
        "academies that had requested those features, and sponsor acceptance of Prototype 3 tax packaging "
        "for Nokta Academy’s compliance workflow. Chapter 4 presents the supporting questionnaire figures.",
    )


def build_chapter4(doc: Document):
    page_break(doc)
    add_h(doc, "CHAPTER 4: RESULTS", 1)
    add_p(
        doc,
        "This chapter presents the results of the questionnaire study and the engineering outcomes of the "
        "three NAMS prototypes. The structure follows Exit.pdf’s results organization—respondent profile, "
        "current practices, feature priorities, perceived impact, and summary—while adapting every measure "
        "to academy management rather than medical exit-exam preparation.",
    )

    add_h(doc, "4.1 Questionnaire Results", 2)
    add_h(doc, "4.1.1 Respondent Profile", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_1_respondent_role.png",
        "4.1",
        "Respondent roles across the 20 surveyed academies",
        [
            "Figure 4.1 shows who answered on behalf of each academy. Administrators and owners dominate, "
            "which is appropriate because they authorize tool adoption and finance workflows.",
            "Teachers and branch managers contribute operational perspective on attendance and exams, while "
            "accountant responses emphasize fee and tax concerns.",
            "Role diversity reduces the risk that requirements reflect only one department’s wishlist "
            "(Creswell & Creswell, 2018).",
        ],
        [
            "The profile in Figure 4.1 supports interpreting later priority scores as institutional rather "
            "than purely student-preference data.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_4_2_academy_maturity.png",
        "4.2",
        "Operating maturity of participating academies",
        [
            "Figure 4.2 separates academies operating more than one year (n=18) from newly started academies "
            "(n=2). This split is analytically important because the two new academies provided no specific "
            "opinion on Prototype 2 enhancements.",
            "Mature academies had experienced spreadsheet pain points long enough to evaluate AI exams, "
            "online class links, and parental attendance automation concretely.",
            "New academies were retained in the sample for honesty about uncertainty rather than forced opinions.",
        ],
        [
            "Figure 4.2 therefore explains why Prototype 2 authorization used an 18/20 demand rule rather "
            "than a unanimous rule.",
        ],
    )

    add_h(doc, "4.1.2 Current Preparation and Access", 3)
    # adapted title: current management tools
    figure_block(
        doc,
        ASSETS / "fig_4_3_current_tools.png",
        "4.3",
        "Current academic–administrative tools used by surveyed academies",
        [
            "Figure 4.3 reports current tools. Most academies still rely on paper/Excel or messaging apps, "
            "with only a minority using partial software. This empirical pattern confirms the problem "
            "statement’s claim of fragmentation.",
            "Messaging-only coordination (for example, chat groups) can distribute announcements quickly "
            "but fails at durable fee balances, audit trails, and tax packaging.",
            "The distribution justifies building an integrated MIS instead of a single-purpose exam app.",
        ],
        [
            "Results in Figure 4.3 also caution implementation: onboarding must import or replace Excel "
            "habits gradually, which again supports phased Agile delivery.",
        ],
    )

    add_h(doc, "4.1.3 Feature Interest and Priorities", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_4_feature_need.png",
        "4.4",
        "Academies requesting Prototype 2 feature package (AI exams, online meeting links, auto parent attendance)",
        [
            "Figure 4.4 is the pivotal requirements result for Prototype 2. Eighteen academies stated that "
            "they need the AI examination, online meeting, and automatic parental attendance package. Two "
            "newly started academies expressed no specific opinion.",
            "In percentage terms, 90% of the sample requested the package, providing clear Product Owner "
            "authority to schedule Prototype 2 after the baseline stabilized.",
            "The figure is not a popularity contest alone; it is tied to interview notes that linked these "
            "features to teacher workload and parent trust.",
        ],
        [
            "Figure 4.4 should be read with Figure 4.5, which ranks a broader feature set by mean priority.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_4_5_priority.png",
        "4.5",
        "Mean priority scores of proposed NAMS feature groups",
        [
            "Figure 4.5 shows mean priority scores on a 1–5 scale. Fee/payroll and attendance lead, followed "
            "closely by exams/AI, parent portal capabilities, tax reports, and messaging.",
            "High scores for fee/payroll validate Prototype 1’s finance spine. High scores for exams/AI and "
            "parent portal validate Prototype 2. Strong tax interest among respondents, combined with "
            "Nokta Academy’s direct request, validates Prototype 3.",
            "The relatively close means indicate that academies want an integrated suite rather than a "
            "single silver-bullet module (DeLone & McLean, 2003).",
        ],
        [
            "Priority results in Figure 4.5 guided sprint ordering when calendar pressure forced sequencing "
            "choices among otherwise important items.",
        ],
    )

    add_h(doc, "4.1.4 Perceived Impact", 3)
    figure_block(
        doc,
        ASSETS / "fig_4_6_impact.png",
        "4.6",
        "Perceived impact statements: number of academies agreeing",
        [
            "Figure 4.6 summarizes agreement with expected impact statements. Nearly all respondents agreed "
            "that NAMS would reduce manual work and improve transparency; most also expected faster reports, "
            "stronger parent trust, and better tax readiness.",
            "Perceived usefulness is a central predictor of acceptance (Davis, 1989; Venkatesh et al., 2003). "
            "These counts therefore support adoption likelihood interpretations in Figure 4.7.",
            "Tax readiness scoring high even before Prototype 3 coding shows latent demand that the sponsor "
            "request later made explicit.",
        ],
        [
            "Impact perceptions in Figure 4.6 are expectations, not post-deployment longitudinal proof; "
            "Chapter 5 treats that distinction carefully.",
        ],
    )
    figure_block(
        doc,
        ASSETS / "fig_4_7_acceptance.png",
        "4.7",
        "Likelihood of adopting NAMS among surveyed academies",
        [
            "Figure 4.7 shows adoption likelihood: 12 academies were very likely, 6 likely, and 2 neutral. "
            "No academy selected “unlikely” in this sample.",
            "Neutrals correspond to the newly started academies that also withheld Prototype 2 opinions, "
            "reinforcing a maturity effect on technology commitment.",
            "Overall acceptance is strong for a purposive sample and supports proceeding from prototype "
            "gates to broader pilot recommendations in Chapter 6.",
        ],
        [
            "Figure 4.7 closes the questionnaire subsection and leads into prototype and engineering results.",
        ],
    )

    add_h(doc, "4.1.5 Summary of Findings", 3)
    add_p(
        doc,
        "Questionnaire findings show fragmented current tools, strong demand for integrated fee/attendance/"
        "exam capabilities, 18/20 demand for Prototype 2 enhancements, high expected impact, and high "
        "adoption likelihood. These findings answer the requirements questions that Chapter 3 framed and "
        "supply evidence for design trade-offs already implemented.",
    )
    add_p(
        doc,
        "Engineering findings complementary to the survey are: Prototype 1 delivered a usable baseline; "
        "Prototype 2 features were implemented with configuration controls; Prototype 3 produced quarterly "
        "and annual tax outputs accepted in sponsor UAT; unit tests protected payroll/tax rules; and "
        "architecture diagrams/sequences remained aligned with the running monorepo modules.",
    )
    add_p(
        doc,
        "Table 4.1 consolidates objective-oriented result statements for quick reference before discussion.",
    )
    table_block(
        doc,
        "4.1",
        "Summary Mapping of Result Evidence to Project Goals",
        ["Goal theme", "Evidence", "Prototype gate"],
        [
            ["Centralize records", "CRUD modules + ERD/schema figures", "P1"],
            ["Improve attendance visibility", "Manual marks (P1); auto online + parent alerts (P2)", "P1–P2"],
            ["Speed exam preparation", "AI generation with teacher review", "P2"],
            ["Support online/hybrid class ops", "Timetable online links + jobs", "P2"],
            ["Financial transparency", "Payments–payroll linkage", "P1"],
            ["Government tax readiness", "Quarterly/annual tax package", "P3"],
            ["Stakeholder acceptance", "18/20 feature demand; 18/20 likely/very likely adopt", "Survey"],
        ],
        [
            "Table 4.1 is a results synthesis, not a new dataset. It exists to help readers transition into "
            "Chapter 5’s interpretive discussion without rereading every figure.",
            "Each row points either to a prototype artifact or to survey evidence presented above.",
        ],
        [
            "The discussion chapter uses Table 4.1 as an index when relating findings to objectives and literature.",
        ],
    )


def build_chapter5(doc: Document):
    page_break(doc)
    add_h(doc, "CHAPTER 5: DISCUSSION", 1)
    add_p(
        doc,
        "This chapter interprets the results. It discusses main findings, relations to objectives and "
        "research questions, relations to literature and existing systems, practical implications for "
        "NAMS, and limitations—following the Exit.pdf discussion pattern while remaining specific to "
        "academy management.",
    )

    add_h(doc, "5.1 Main Findings", 2)
    add_p(
        doc,
        "Three findings dominate. First, academies still operate with fragmented tools, so an integrated "
        "MIS addresses a real operational gap. Second, feature demand is broad but not uniform over time: "
        "core finance/academic modules were needed immediately, while AI exams, online meeting links, and "
        "parental auto-attendance were justified after 18/20 academies requested them. Third, compliance "
        "reporting (Prototype 3) emerged as a sponsor-critical capability that generic LMS platforms rarely "
        "treat as native (Rahman & Saha, 2019).",
    )
    add_p(
        doc,
        "These findings jointly support an iterative product strategy. A single big-bang release attempting "
        "AI, automation, and tax in the first demo would have delayed the baseline that staff needed most "
        "(Dikert et al., 2016; Pressman & Maxim, 2020).",
    )
    add_p(
        doc,
        "Acceptance likelihood was high, but neutrality among new academies shows that digital readiness "
        "moderates adoption (Venkatesh et al., 2003).",
    )

    add_h(doc, "5.2 Relation to Objectives and Research Questions", 2)
    add_p(
        doc,
        "Chapter 1 objectives sought centralized information management, better communication, academic "
        "monitoring, operational efficiency, and analytics readiness. Prototype 1 advanced centralization "
        "and efficiency; Prototype 2 advanced communication/monitoring for online contexts and exam "
        "authoring; Prototype 3 advanced institutional reporting for tax compliance. Analytics readiness "
        "is partially achieved through structured data capture even where advanced ML models are future work "
        "(Aldowah et al., 2019).",
    )
    add_p(
        doc,
        "Research questions about fragmentation, timely information, communication, analytics, and "
        "automation are therefore answered at the level of an engineered and reviewed system plus "
        "stakeholder evidence, not as a multi-year causal impact study. That scope matches a bachelor "
        "final-year project while still producing actionable knowledge (Hevner et al., 2004).",
    )
    add_p(
        doc,
        "Table 5.1 provides an explicit answer matrix.",
    )
    table_block(
        doc,
        "5.1",
        "Discussion Matrix Linking Research Questions to Evidence",
        ["Question theme", "Answer from NAMS evidence"],
        [
            ["Fragmented records", "Integrated student–class–fee–attendance–exam data model reduces multi-file reconciliation."],
            ["Timely integrated information", "Dashboards and role portals expose current balances, attendance, and results."],
            ["Communication/coordination", "Announcements/messages plus Prototype 2 parent attendance alerts shorten informal chasing."],
            ["Monitoring/analytics", "Structured histories enable reports now and AI/analytics extensions later."],
            ["Automation efficiency", "Fee/payroll linkage and Prototype 2/3 automations cut repetitive manual steps."],
        ],
        [
            "Table 5.1 compresses discussion claims into audit-friendly statements for defense and grading.",
            "Each answer is bounded by the sample and prototype scope described in Section 5.5.",
        ],
        [
            "The matrix in Table 5.1 also shows why methodology choices (mixed methods + Agile prototypes) "
            "were not cosmetic: they produced the evidence types each answer requires.",
        ],
    )

    add_h(doc, "5.3 Relation to Literature and Existing Systems", 2)
    add_p(
        doc,
        "Relative to Moodle, Google Classroom, and Canvas, NAMS is less of a content LMS and more of an "
        "academy operations platform. Literature already notes that instructional platforms often need "
        "external tools for finance and institutional reporting (Costa et al., 2020). NAMS inverts that "
        "pattern by centering operations and attaching learning-support features where they help academies.",
    )
    add_p(
        doc,
        "Information systems success theory helps interpret why tax and payroll accuracy matter: information "
        "quality and net benefits are damaged when financial figures are doubted (DeLone & McLean, 2003). "
        "Technology acceptance theory helps interpret why Prototype 2 automation had to remain reviewable "
        "by teachers—perceived control affects trust (Davis, 1989).",
    )
    add_p(
        doc,
        "Learning analytics literature supports storing attendance and assessment consistently even before "
        "advanced models ship (Romero & Ventura, 2020; Siemens, 2013). NAMS follows that sequencing.",
    )

    add_h(doc, "5.4 Practical Implications of NAMS", 2)
    add_p(
        doc,
        "Practically, NAMS offers private academies a path from Excel/chat coordination to a role-aware "
        "system that can grow. Prototype gating reduces adoption shock: an academy may begin with Prototype "
        "1 features and enable Prototype 2 flags when teachers and parents are ready.",
    )
    add_p(
        doc,
        "For Nokta Academy specifically, Prototype 3 reduces dependency on ad-hoc tax spreadsheets and "
        "creates an auditable quarterly/annual trail. For other academies in the survey, the 18/20 demand "
        "signal suggests a transferable enhancement package around AI exams, online links, and parental "
        "attendance visibility.",
    )
    add_p(
        doc,
        "Multilingual and offline-friendly deployment considerations remain practical differentiators in "
        "the local context discussed in Chapter 2.",
    )

    add_h(doc, "5.5 Limitations of the Findings and System", 2)
    add_p(
        doc,
        "Limitations include purposive sampling of 20 academies, self-reported expectations rather than "
        "longitudinal outcome measurement, and prototype UI figures that include engineered mock "
        "composites where live screenshots were insufficient for print clarity. AI question quality depends "
        "on external providers and teacher review discipline.",
    )
    add_p(
        doc,
        "Government tax rules are implemented according to sponsor-provided operational interpretation for "
        "Nokta Academy; legal advice beyond that encoding is outside the software engineering scope. "
        "Multi-campus national integrations and biometric attendance were out of scope as stated in Chapter 1.",
    )
    add_p(
        doc,
        "These limitations do not void the contribution; they bound claims appropriately for an applied "
        "bachelor project (Hevner et al., 2004; Creswell & Creswell, 2018).",
    )


def build_chapter6(doc: Document):
    page_break(doc)
    add_h(doc, "CHAPTER 6: CONCLUSION AND FUTURE WORK", 1)

    add_h(doc, "6.1 Summary of Work Done", 2)
    add_p(
        doc,
        "This project designed and implemented the Nokta Academy Management System as an integrated "
        "web–API–mobile platform for private academy operations. Chapters 1–2 established the problem and "
        "literature context. Chapter 3 documented mixed-methods requirements work and Agile Scrum delivery "
        "through three prototypes. Chapters 4–5 presented and interpreted survey and engineering results.",
    )
    add_p(
        doc,
        "Prototype 1 delivered the operational baseline without AI exam generation, without online-meeting "
        "attendance automation, and without full government tax packaging. Prototype 2 added those "
        "academic/parent enhancements after 18 of 20 academies requested them. Prototype 3 added quarterly "
        "and annual government tax reporting at Nokta Academy’s request.",
    )
    add_p(
        doc,
        "Design artifacts (use cases, sequences, architecture, ERD/schema) and tests were maintained to "
        "keep the monograph aligned with the running system.",
    )

    add_h(doc, "6.2 Key Contributions", 2)
    add_p(
        doc,
        "Contributions include: (1) an integrated academy MIS spanning academic and financial operations; "
        "(2) an evidence-based prototype strategy linking a 20-academy questionnaire to backlog decisions; "
        "(3) parental online-attendance visibility with policy controls; (4) AI-assisted exam authoring "
        "with human approval; and (5) a quarterly/annual tax reporting package grounded in operational data.",
    )
    add_p(
        doc,
        "Methodologically, the project shows how Agile Scrum practices (backlog, sprints, reviews, "
        "definition of done) can be combined with formal requirements and UML-style design documentation "
        "suitable for academic assessment (Schwaber & Sutherland, 2020; Fowler, 2004).",
    )
    add_p(
        doc,
        "Empirically, the 18/20 feature-demand result provides a clear, citable decision record for why "
        "Prototype 2 existed as a separate gate rather than as hidden incomplete work.",
    )

    add_h(doc, "6.3 Achievement of Objectives", 2)
    table_block(
        doc,
        "6.1",
        "Objective Achievement Status for NAMS",
        ["Objective theme", "Status", "Evidence"],
        [
            ["Centralized information management", "Achieved", "Integrated modules + data model"],
            ["Improved communication/coordination", "Achieved (enhanced in P2)", "Messages + parent attendance alerts"],
            ["Academic monitoring & results access", "Achieved", "Exams/results portals"],
            ["Operational efficiency (attendance/fees/exams)", "Achieved", "P1 baseline + automation in P2"],
            ["Analytics/AI readiness", "Partially achieved", "Structured data + AI authoring; predictive ML future"],
            ["Government tax reporting", "Achieved (P3)", "Quarterly/annual tax package UAT"],
        ],
        [
            "Table 6.1 uses conservative status labels. “Partially achieved” for advanced analytics avoids "
            "over-claiming machine-learning interventions not delivered as production predictors "
            "(Peffers et al., 2007).",
            "Tax reporting is marked achieved for the sponsor-defined package reviewed in Prototype 3 UAT.",
        ],
        [
            "Overall, Table 6.1 indicates that the project met its engineering objectives within the stated "
            "scope and limitations.",
        ],
    )

    add_h(doc, "6.4 Future Enhancements", 2)
    add_p(
        doc,
        "Future work should include longitudinal pilots measuring time saved and error reduction after "
        "deployment; deeper learning-analytics models for at-risk prediction; richer offline "
        "synchronization; expanded Flutter feature parity; stronger backup/recovery drills; and formal "
        "accessibility audits. Payment-gateway and biometric options may be reconsidered if academies "
        "fund those integrations.",
    )
    add_p(
        doc,
        "On the compliance side, tax rule packs could become configurable per jurisdiction if NAMS expands "
        "beyond the current sponsor interpretation. On the AI side, item-quality scoring and curriculum "
        "alignment checks would improve teacher trust in generated questions.",
    )
    add_p(
        doc,
        "Finally, a larger multi-city academy sample would strengthen generalizability beyond the purposive "
        "20-academy questionnaire used here (Creswell & Creswell, 2018).",
    )


def build_references(doc: Document):
    page_break(doc)
    add_h(doc, "REFERENCES", 1)
    refs = [
        "Aldowah, H., Al-Samarraie, H., & Fauzy, W. M. (2019). Educational data mining and learning analytics for 21st century higher education: A review and synthesis. Telematics and Informatics, 37, 13–49. https://doi.org/10.1016/j.tele.2019.01.007",
        "American Psychological Association. (2020). Publication manual of the American Psychological Association (7th ed.). American Psychological Association.",
        "Bass, L., Clements, P., & Kazman, R. (2012). Software architecture in practice (3rd ed.). Addison-Wesley.",
        "Brooke, J. (1996). SUS: A quick and dirty usability scale. In P. W. Jordan, B. Thomas, B. A. Weerdmeester, & I. L. McClelland (Eds.), Usability evaluation in industry (pp. 189–194). Taylor & Francis.",
        "Chodorow, K. (2013). MongoDB: The definitive guide (2nd ed.). O’Reilly Media.",
        "Connolly, T. M., & Begg, C. E. (2015). Database systems: A practical approach to design, implementation, and management (6th ed.). Pearson.",
        "Costa, C., Teixeira, L., & Alvelos, H. (2020). Exploring the usage of learning management systems in higher education: A systematic review. Education and Information Technologies, 25, 1–20.",
        "Creswell, J. W., & Creswell, J. D. (2018). Research design: Qualitative, quantitative, and mixed methods approaches (5th ed.). SAGE Publications.",
        "Davis, F. D. (1989). Perceived usefulness, perceived ease of use, and user acceptance of information technology. MIS Quarterly, 13(3), 319–340. https://doi.org/10.2307/249008",
        "DeLone, W. H., & McLean, E. R. (2003). The DeLone and McLean model of information systems success: A ten-year update. Journal of Management Information Systems, 19(4), 9–30. https://doi.org/10.1080/07421222.2003.11045748",
        "Dikert, K., Paasivaara, M., & Lassenius, C. (2016). Challenges and success factors for large-scale agile transformations: A systematic literature review. Journal of Systems and Software, 119, 87–108. https://doi.org/10.1016/j.jss.2016.06.013",
        "Fowler, M. (2004). UML distilled: A brief guide to the standard object modeling language (3rd ed.). Addison-Wesley.",
        "Hevner, A. R., March, S. T., Park, J., & Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75–105.",
        "International Organization for Standardization, & International Electrotechnical Commission. (2011). Systems and software engineering — Systems and software Quality Requirements and Evaluation (SQuaRE) — System and software quality models (ISO/IEC 25010:2011).",
        "International Organization for Standardization, International Electrotechnical Commission, & Institute of Electrical and Electronics Engineers. (2018). Systems and software engineering — Life cycle processes — Requirements engineering (ISO/IEC/IEEE 29148:2018).",
        "Likert, R. (1932). A technique for the measurement of attitudes. Archives of Psychology, 22(140), 1–55.",
        "Miles, M. B., Huberman, A. M., & Saldaña, J. (2014). Qualitative data analysis: A methods sourcebook (3rd ed.). SAGE Publications.",
        "Nielsen, J. (1994). Usability engineering. Morgan Kaufmann.",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45–77.",
        "Pressman, R. S., & Maxim, B. R. (2020). Software engineering: A practitioner’s approach (9th ed.). McGraw-Hill Education.",
        "Rahman, M. M., & Saha, S. (2019). An integrated education management information system for private academic institutions. International Journal of Information Engineering and Electronic Business, 11(2), 1–10.",
        "Romero, C., & Ventura, S. (2020). Educational data mining and learning analytics: An updated survey. WIREs Data Mining and Knowledge Discovery, 10(3), e1355. https://doi.org/10.1002/widm.1355",
        "Schwaber, K., & Sutherland, J. (2020). The Scrum guide: The definitive guide to Scrum: The rules of the game. Scrum.org.",
        "Siemens, G. (2013). Learning analytics: The emergence of a discipline. American Behavioral Scientist, 57(10), 1380–1400. https://doi.org/10.1177/0002764213498851",
        "Sommerville, I. (2016). Software engineering (10th ed.). Pearson.",
        "Venkatesh, V., Morris, M. G., Davis, G. B., & Davis, F. D. (2003). User acceptance of information technology: Toward a unified view. MIS Quarterly, 27(3), 425–478.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        run = p.add_run(ref)
        set_run(run, size=BODY_PT)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.first_line_indent = Inches(-0.5)
        pf.left_indent = Inches(0.5)
        pf.space_after = Pt(8)


def build_appendices(doc: Document):
    page_break(doc)
    add_h(doc, "APPENDIX A: QUESTIONNAIRES", 1)
    add_h(doc, "A.1 Requirements Elicitation Questionnaire", 2)
    add_p(
        doc,
        "The questionnaire was administered to one primary respondent per academy (n=20). Items used "
        "categorical choices and five-point Likert-type ratings (Likert, 1932). The instrument is summarized "
        "below; full printed forms were used during data collection.",
        first_indent=True,
    )
    add_p(doc, "Section A — Academy profile: academy age, approximate student count, respondent role.", first_indent=True)
    add_p(doc, "Section B — Current tools: paper/Excel, messaging apps, partial software, none.", first_indent=True)
    add_p(doc, "Section C — Pain points: enrollment, attendance, exams, fees, parent communication, reporting, tax.", first_indent=True)
    add_p(doc, "Section D — Feature usefulness and priority for core MIS modules (1–5).", first_indent=True)
    add_p(
        doc,
        "Section E — Prototype 2 package: need for AI exam support, online class meeting links, and automatic "
        "parental attendance visibility (Yes / No specific opinion).",
        first_indent=True,
    )
    add_p(doc, "Section F — Perceived impact statements and adoption likelihood.", first_indent=True)
    add_p(
        doc,
        "Scoring note: eighteen academies answered “Yes” to Section E; two newly started academies selected "
        "“No specific opinion.” These counts authorized Prototype 2 scope in Chapter 3.",
        first_indent=True,
    )

    page_break(doc)
    add_h(doc, "APPENDIX B: SOURCE CODE REPOSITORY", 1)
    add_p(
        doc,
        "The NAMS source code is maintained as a monorepo containing backend (Node.js/Express/TypeScript), "
        "web frontend (React/Vite), mobile/PWA assets, deployment configuration, ERD documentation, and "
        "monograph generation scripts. Local development uses MongoDB and environment-variable configuration "
        "described in the repository README.",
        first_indent=True,
    )
    add_p(
        doc,
        "Key paths include backend/src for API modules and jobs, frontend/src/features for UI domains, "
        "ERD/ for schema documentation, and scripts/ for report generation utilities used to produce this document.",
        first_indent=True,
    )
    add_p(
        doc,
        "Access credentials for demonstration accounts, if required by examiners, are provided separately "
        "through the department’s project submission channel and are not printed in the public monograph body.",
        first_indent=True,
    )

    page_break(doc)
    add_h(doc, "APPENDIX C: RAW TEST DATA", 1)
    add_p(
        doc,
        "Raw automated test outputs and CI logs are retained in the project repository. This appendix "
        "summarizes representative cases already presented in Table 3.6 and confirms that payroll, tax, "
        "RBAC, and attendance-policy tests were executed before prototype gate reviews.",
        first_indent=True,
    )
    table_block(
        doc,
        "C.1",
        "Representative Raw Test Outcome Summary",
        ["Suite", "Focus", "Outcome at Prototype-3 gate"],
        [
            ["Finance unit tests", "Payroll percentage + tax brackets", "Pass"],
            ["Academy tax tests", "Quarterly BRT aggregation", "Pass"],
            ["Attendance policy tests", "Auto-mark on/off behavior", "Pass"],
            ["Auth/RBAC tests", "Role denial on sensitive routes", "Pass"],
            ["Exam/AI smoke tests", "Generate–edit–publish path", "Pass"],
        ],
        [
            "Table C.1 is a compact raw-outcome digest for examiners. Detailed assertion messages remain in "
            "the repository test runner output.",
            "A “Pass” status means the suite succeeded on the evaluated commit used for Prototype 3 UAT.",
        ],
        [
            "If department policy requires attaching full console logs, they can be exported from CI without "
            "altering the narrative chapters.",
        ],
    )

    page_break(doc)
    add_h(doc, "APPENDIX D: USER MANUAL", 1)
    add_p(
        doc,
        "This abbreviated user manual describes common tasks for NAMS roles. It is intended for examiner "
        "orientation and academy onboarding, not as a substitute for in-app help text.",
        first_indent=True,
    )
    add_h(doc, "D.1 Administrator", 2)
    add_p(doc, "1) Sign in with an admin account. 2) Create branches/classes/subjects as needed. 3) Register students and assign teachers. 4) Record fees and review dashboards. 5) For Prototype 3, open Payroll/Tax reports, select Hijri quarter or year, and export PDF.", first_indent=True)
    add_h(doc, "D.2 Teacher", 2)
    add_p(doc, "1) Open assigned classes. 2) Mark attendance for the session. 3) Create exams; optionally generate AI draft questions (Prototype 2), edit, and publish. 4) Enter or review results. 5) Check timetable online links for hybrid sessions.", first_indent=True)
    add_h(doc, "D.3 Student", 2)
    add_p(doc, "1) Sign in to view timetable, attendance summary, results, fee status, and announcements.", first_indent=True)
    add_h(doc, "D.4 Parent", 2)
    add_p(doc, "1) Sign in to the family/parent portal. 2) View linked children’s attendance and results. 3) In Prototype 2, receive notifications when online auto-attendance records are created for enabled sessions.", first_indent=True)
    add_h(doc, "D.5 Owner / Super Admin", 2)
    add_p(doc, "1) Review academy-level dashboards. 2) Manage roles/permissions as authorized. 3) Export quarterly/annual tax packages and archive them for compliance review.", first_indent=True)
    add_p(
        doc,
        "End of monograph appendices.",
        first_indent=False,
        align="center",
    )


def main():
    src = find_source_docx()
    print("SOURCE", str(src).encode("unicode_escape").decode())
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(src, OUT)
    doc = Document(str(OUT))
    configure_styles(doc)

    # Ensure required asset files exist
    required = [
        ASSETS / "fig_agile_scrum.png",
        ASSETS / "fig_3_1_proto1_dashboard.png",
        ASSETS / "fig_3_12_use_case.png",
        ASSETS / "fig_4_1_respondent_role.png",
        MONO / "Figure_3_13_Student_Enrollment_and_Fee_Payment_Sequence.png",
        MONO / "Figure_3_16_NAMS_Architecture_Diagram.png",
        MONO / "Figure_3_17_NAMS_Entity_Relationship_Diagram.png",
        MONO / "Figure_3_18_NAMS_Core_Relational_Document_Schema.png",
    ]
    missing = [str(p) for p in required if not p.exists()]
    if missing:
        raise FileNotFoundError("Missing assets:\n" + "\n".join(missing))

    build_chapter3(doc)
    build_chapter4(doc)
    build_chapter5(doc)
    build_chapter6(doc)
    build_references(doc)
    build_appendices(doc)

    doc.save(str(OUT))
    size_mb = OUT.stat().st_size / (1024 * 1024)
    print("CREATED", str(OUT).encode("unicode_escape").decode())
    print("SIZE_MB", round(size_mb, 2))


if __name__ == "__main__":
    main()
